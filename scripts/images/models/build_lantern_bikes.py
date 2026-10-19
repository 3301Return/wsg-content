#!/usr/bin/env python3
"""Build scripts/images/models/lantern-bikes.xlsx for the WSG article
"The Three Financial Statements: How They Link and How to Explain It" (October 5, 2026).

Three sheets, every computed cell a live formula, inputs blue on pale yellow:

  "Lantern Bikes"       section 6: the Year 2 drivers, the Year 2 income statement, the Year 2
                        cash flow statement and the Year 1 to Year 2 balance sheet with the
                        balance check row. Column G carries the article's "where it goes" notes,
                        column I the Year 2 formula as text for the close-ups.
  "Depreciation walk"   section 4: the $10 of depreciation walk at a 25 percent tax rate, with
                        the 40 percent column older guides print, formulas in column G.
  "Scenarios"           section 5: the ten scenario tables, one block each, laid out exactly as
                        the article prints them (inputs on the first row or two of each block).

The numbers are the article's own and also come out of
scripts/images/models/numbers_three-financial-statements-how-they-link.py, which this script
imports to cross-check the Lantern Bikes and depreciation figures. Every formula cell is
written together with the value the article's table shows for it, and check_against_article()
recalculates the workbook with LibreOffice and compares cell for cell.

Run from the repo root:  python3 scripts/images/models/build_lantern_bikes.py
"""
import importlib.util
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts" / "images"))
from wsg_viz import (FILL_SUBTLE, XL_FONT, XL_SIZE, header_bar, new_workbook, put,  # noqa: E402
                     recalc_values, set_widths)

OUT = ROOT / "scripts" / "images" / "models" / "lantern-bikes.xlsx"
NUMBERS = ROOT / "scripts" / "images" / "models" / "numbers_three-financial-statements-how-they-link.py"

FMT_INT = '#,##0;(#,##0);0'            # the Lantern Bikes tables print whole numbers, zero as 0
FMT_DEC = '#,##0.0;(#,##0.0);0.0'      # the walks print 2.5, 7.5, 12.5, 37.5
FMT_PCT0 = '0%'
GRAY = "52514E"
MUTED = "8A8984"

# {sheet title: {cell: value the article's table shows}} filled while the sheets are written
EXPECTED = {}


def load_numbers():
    """Import the drafting agent's numbers script (its file name has hyphens)."""
    spec = importlib.util.spec_from_file_location("wsg_numbers", NUMBERS)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------- small helpers
def expect(ws, ref, value):
    EXPECTED.setdefault(ws.title, {})[ref] = value


def title(ws, ref, text):
    put(ws, ref, text, bold=True, size=12)


def note(ws, ref, text):
    put(ws, ref, text, italic=True, color=GRAY, size=9)


def formula_text(ws, ref, text):
    """The formula written out as text beside the row so a screenshot can show the link."""
    c = put(ws, ref, text, color=GRAY, align="left", indent=1)
    c.data_type = "s"
    return c


NUMERIC_HEADERS = {"Year 1", "Year 2", "Change", "25% tax rate", "40% tax rate"}


def section(ws, row, text, cols, last_col=9):
    """A navy bar from B to last_col with white column labels in the cells named by cols
    (number columns right-aligned, text columns such as Formula or Why left-aligned)."""
    from openpyxl.styles import Alignment, Font
    header_bar(ws, row, 2, last_col, text)
    for col, label in cols.items():
        c = ws[f"{col}{row}"]
        c.value = label
        c.font = Font(name=XL_FONT, size=XL_SIZE, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="right" if label in NUMERIC_HEADERS else "left",
                                indent=0 if label in NUMERIC_HEADERS else 1)


def cell(ws, ref, formula, want, kind="formula", fmt=FMT_INT):
    """A formula cell plus the value the article prints for it."""
    put(ws, ref, formula, kind=kind, fmt=fmt)
    expect(ws, ref, want)


def no_change(ws, ref):
    put(ws, ref, "no change", color=MUTED, align="right")


# ---------------------------------------------------------------- sheet 1: Lantern Bikes
R = dict(
    drv_hdr=4, rev_in=5, cogs_in=6, sga_in=7, da_in=8, rate_in=9, tax_in=10, capex_in=11,
    ar_in=12, inv_in=13, ap_in=14, acc_in=15, debt_in=16, div_in=17,
    is_hdr=19, rev=20, cogs=21, gp=22, sga=23, da=24, ebit=25, interest=26, ebt=27, taxes=28, ni=29,
    cf_hdr=31, cf_ni=32, cf_da=33, cf_ar=34, cf_inv=35, cf_ap=36, cf_acc=37, cfo=38, cf_capex=39,
    cfi=40, cf_debt=41, cf_div=42, cff=43, netchg=44, begcash=45, endcash=46,
    bs_hdr=48, cash=49, ar=50, inv=51, ppe=52, ta=53, ap=54, acc=55, debt=56, tl=57, cs=58, re=59,
    te=60, tle=61, check=62,
)


def build_lantern(wb, numbers):
    ws = wb.active
    ws.title = "Lantern Bikes"
    set_widths(ws, {"A": 2, "B": 48, "C": 2, "D": 10, "E": 10, "F": 10, "G": 48, "H": 2, "I": 24})
    ws.sheet_view.showGridLines = False
    title(ws, "B1", "Lantern Bikes Co. (fictional), Year 1 to Year 2")
    note(ws, "B2", "$ millions. Blue cells on yellow are the article's inputs; every other number is a formula. "
                   "Parentheses mean a negative number.")

    base = numbers.lantern_bikes(print_tables=False)   # the drafting script's own figures
    y2 = base["y2"]

    # ---- Year 2 drivers (inputs)
    section(ws, R["drv_hdr"], "Year 2 drivers", {"E": "Year 2"}, last_col=7)
    drivers = [
        ("rev_in", "Revenue", 500, FMT_INT), ("cogs_in", "Cost of goods sold", 300, FMT_INT),
        ("sga_in", "SG&A", 100, FMT_INT), ("da_in", "Depreciation and amortization", 30, FMT_INT),
        ("rate_in", "Interest rate on beginning debt", 0.05, FMT_PCT0), ("tax_in", "Tax rate", 0.25, FMT_PCT0),
        ("capex_in", "Capital expenditures", 40, FMT_INT),
        ("ar_in", "Increase in accounts receivable", 10, FMT_INT), ("inv_in", "Increase in inventory", 5, FMT_INT),
        ("ap_in", "Increase in accounts payable", 6, FMT_INT), ("acc_in", "Increase in accrued expenses", 2, FMT_INT),
        ("debt_in", "Debt repaid", 20, FMT_INT), ("div_in", "Dividends paid", 12, FMT_INT),
    ]
    for key, label, val, fmt in drivers:
        put(ws, f"B{R[key]}", label, indent=1)
        put(ws, f"E{R[key]}", val, kind="input", fmt=fmt)

    # ---- Year 2 income statement
    section(ws, R["is_hdr"], "Income statement", {"E": "Year 2", "G": "Where it goes", "I": "Formula (Year 2)"})
    is_rows = [
        # key, label, formula, article value, kind, note
        ("rev", "Revenue", f"=E{R['rev_in']}", 500, "formula", ""),
        ("cogs", "Cost of goods sold", f"=-E{R['cogs_in']}", -300, "formula", ""),
        ("gp", "Gross profit", f"=E{R['rev']}+E{R['cogs']}", 200, "total", ""),
        ("sga", "SG&A", f"=-E{R['sga_in']}", -100, "formula", ""),
        ("da", "Depreciation and amortization", f"=-E{R['da_in']}", -30, "formula",
         "Link 2: added back on the CFS, reduces PP&E"),
        ("ebit", "EBIT (operating profit)", f"=SUM(E{R['gp']}:E{R['da']})", 70, "total", ""),
        ("interest", "Interest expense (5% on beginning debt of 200)", f"=-E{R['rate_in']}*D{R['debt']}", -10,
         "formula", "Computed off the Year 1 debt balance"),
        ("ebt", "Pre-tax income", f"=E{R['ebit']}+E{R['interest']}", 60, "total", ""),
        ("taxes", "Taxes (25%)", f"=-E{R['ebt']}*E{R['tax_in']}", -15, "formula", ""),
        ("ni", "Net income", f"=E{R['ebt']}+E{R['taxes']}", 45, "total",
         "Link 1: top of the CFS and into retained earnings"),
    ]
    for key, label, fml, want, kind, why in is_rows:
        r = R[key]
        put(ws, f"B{r}", label, bold=(kind == "total"), indent=0 if kind == "total" else 1)
        cell(ws, f"E{r}", fml, want, kind=kind)
        if why:
            put(ws, f"G{r}", why, color=GRAY, align="left", indent=1)
        formula_text(ws, f"I{r}", fml)

    # ---- Year 2 cash flow statement
    section(ws, R["cf_hdr"], "Cash flow statement", {"E": "Year 2", "G": "Where it comes from or goes",
                                                      "I": "Formula (Year 2)"})
    cf_rows = [
        ("cf_ni", "Net income", f"=E{R['ni']}", 45, "formula", "Link 1: from the income statement"),
        ("cf_da", "Plus: depreciation and amortization", f"=-E{R['da']}", 30, "formula", "Link 2: non-cash, added back"),
        ("cf_ar", "Less: increase in accounts receivable", f"=-(E{R['ar']}-D{R['ar']})", -10, "formula",
         "Link 4: balance sheet AR 60 to 70"),
        ("cf_inv", "Less: increase in inventory", f"=-(E{R['inv']}-D{R['inv']})", -5, "formula",
         "Link 4: inventory 80 to 85"),
        ("cf_ap", "Plus: increase in accounts payable", f"=E{R['ap']}-D{R['ap']}", 6, "formula",
         "Link 4: payables 50 to 56"),
        ("cf_acc", "Plus: increase in accrued expenses", f"=E{R['acc']}-D{R['acc']}", 2, "formula",
         "Link 4: accrued 20 to 22"),
        ("cfo", "Cash from operations", f"=SUM(E{R['cf_ni']}:E{R['cf_acc']})", 68, "total", ""),
        ("cf_capex", "Capital expenditures", f"=-E{R['capex_in']}", -40, "formula", "Link 3: raises PP&E"),
        ("cfi", "Cash from investing", f"=E{R['cf_capex']}", -40, "total", ""),
        ("cf_debt", "Debt repaid", f"=E{R['debt']}-D{R['debt']}", -20, "formula", "Link 5: debt 200 to 180"),
        ("cf_div", "Dividends paid", f"=-E{R['div_in']}", -12, "formula", "Reduces retained earnings"),
        ("cff", "Cash from financing", f"=E{R['cf_debt']}+E{R['cf_div']}", -32, "total", ""),
        ("netchg", "Net change in cash", f"=E{R['cfo']}+E{R['cfi']}+E{R['cff']}", -4, "total", ""),
        ("begcash", "Beginning cash", f"=D{R['cash']}", 40, "formula", "Year 1 balance sheet"),
        ("endcash", "Ending cash", f"=E{R['begcash']}+E{R['netchg']}", 36, "total",
         "Link 6: becomes cash on the Year 2 balance sheet"),
    ]
    for key, label, fml, want, kind, why in cf_rows:
        r = R[key]
        put(ws, f"B{r}", label, bold=(kind == "total"), indent=0 if kind == "total" else 1)
        cell(ws, f"E{r}", fml, want, kind=kind)
        if why:
            put(ws, f"G{r}", why, color=GRAY, align="left", indent=1)
        formula_text(ws, f"I{r}", fml)

    # ---- Balance sheet, Year 1 (inputs) to Year 2 (formulas)
    section(ws, R["bs_hdr"], "Balance sheet", {"D": "Year 1", "E": "Year 2", "F": "Change", "G": "Why",
                                                "I": "Formula (Year 2)"})
    bs_rows = [
        # key, label, Year 1 input or formula, Year 2 formula, article (Y1, Y2, change), kind, why
        ("cash", "Cash", 40, f"=E{R['endcash']}", (40, 36, -4), "formula", "Link 6: ending cash from the CFS"),
        ("ar", "Accounts receivable", 60, f"=D{R['ar']}+E{R['ar_in']}", (60, 70, 10), "formula", "Link 4"),
        ("inv", "Inventory", 80, f"=D{R['inv']}+E{R['inv_in']}", (80, 85, 5), "formula", "Link 4"),
        ("ppe", "Net PP&E", 300, f"=D{R['ppe']}+E{R['capex_in']}-E{R['da_in']}", (300, 310, 10), "formula",
         "Links 2 and 3: plus capex 40, minus D&A 30"),
        ("ta", "Total assets", f"=SUM(D{R['cash']}:D{R['ppe']})", f"=SUM(E{R['cash']}:E{R['ppe']})",
         (480, 501, 21), "total", ""),
        ("ap", "Accounts payable", 50, f"=D{R['ap']}+E{R['ap_in']}", (50, 56, 6), "formula", "Link 4"),
        ("acc", "Accrued expenses", 20, f"=D{R['acc']}+E{R['acc_in']}", (20, 22, 2), "formula", "Link 4"),
        ("debt", "Long-term debt", 200, f"=D{R['debt']}-E{R['debt_in']}", (200, 180, -20), "formula",
         "Link 5: repaid 20"),
        ("tl", "Total liabilities", f"=SUM(D{R['ap']}:D{R['debt']})", f"=SUM(E{R['ap']}:E{R['debt']})",
         (270, 258, -12), "total", ""),
        ("cs", "Common stock", 100, f"=D{R['cs']}", (100, 100, 0), "formula", "No shares issued or bought back"),
        ("re", "Retained earnings", 110, f"=D{R['re']}+E{R['ni']}-E{R['div_in']}", (110, 143, 33), "formula",
         "Link 1: plus net income 45, minus dividends 12"),
        ("te", "Total equity", f"=D{R['cs']}+D{R['re']}", f"=E{R['cs']}+E{R['re']}", (210, 243, 33), "total", ""),
        ("tle", "Total liabilities and equity", f"=D{R['tl']}+D{R['te']}", f"=E{R['tl']}+E{R['te']}",
         (480, 501, 21), "total", ""),
    ]
    for key, label, y1, fml, (w1, w2, wc), kind, why in bs_rows:
        r = R[key]
        put(ws, f"B{r}", label, bold=(kind == "total"), indent=0 if kind == "total" else 1)
        if isinstance(y1, str):
            cell(ws, f"D{r}", y1, w1, kind=kind)
        else:
            put(ws, f"D{r}", y1, kind="input", fmt=FMT_INT)
            expect(ws, f"D{r}", w1)
        cell(ws, f"E{r}", fml, w2, kind=kind)
        cell(ws, f"F{r}", f"=E{r}-D{r}", wc, kind=kind)
        if why:
            put(ws, f"G{r}", why, color=GRAY, align="left", indent=1)
        formula_text(ws, f"I{r}", fml)
    r = R["check"]
    put(ws, f"B{r}", "Balance check (assets minus L+E)", bold=True)
    cell(ws, f"D{r}", f"=D{R['ta']}-D{R['tle']}", 0)
    cell(ws, f"E{r}", f"=E{R['ta']}-E{R['tle']}", 0)
    put(ws, f"G{r}", "Must be 0", color=GRAY, align="left", indent=1)
    formula_text(ws, f"I{r}", f"=E{R['ta']}-E{R['tle']}")
    from openpyxl.styles import PatternFill
    for col in "BDEFG":
        ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=FILL_SUBTLE)

    # cross-check the article values typed above against the drafting script's own output
    assert (base["ni"], base["cfo"], base["cfi"], base["cff"], base["net_change"], base["ending_cash"]) == \
        (45, 68, -40, -32, -4, 36)
    assert (y2["cash"], y2["ar"], y2["inventory"], y2["ppe"], y2["ap"], y2["accrued"], y2["debt"],
            y2["common"], y2["re"]) == (36, 70, 85, 310, 56, 22, 180, 100, 143)
    assert base["check"] == 0
    return ws


# ---------------------------------------------------------------- sheet 2: the $10 depreciation walk
def build_depreciation(wb, numbers):
    ws = wb.create_sheet("Depreciation walk")
    set_widths(ws, {"A": 2, "B": 50, "C": 2, "D": 13, "E": 13, "F": 2, "G": 22})
    ws.sheet_view.showGridLines = False
    title(ws, "B1", "Walk me through $10 of depreciation")
    note(ws, "B2", "$. Each line is the change versus a year without the extra depreciation. "
                   "Parentheses mean the line went down.")

    section(ws, 4, "Inputs", {"D": "25% tax rate", "E": "40% tax rate"}, last_col=7)
    put(ws, "B5", "Extra depreciation", indent=1)
    put(ws, "B6", "Tax rate", indent=1)
    put(ws, "D5", 10, kind="input", fmt=FMT_DEC)
    put(ws, "E5", 10, kind="input", fmt=FMT_DEC)
    put(ws, "D6", 0.25, kind="input", fmt=FMT_PCT0)
    put(ws, "E6", 0.40, kind="input", fmt=FMT_PCT0)

    w25 = numbers.depreciation_walk(10, 0.25)
    w40 = numbers.depreciation_walk(10, 0.40)
    blocks = [
        (8, "Income statement", [
            ("Depreciation expense", "={c}5", 10, 10, "formula"),
            ("Pre-tax income", "=-{c}9", -10, -10, "formula"),
            ("Taxes", "={c}10*{c}6", -2.5, -4, "formula"),
            ("Net income", "={c}10-{c}11", -7.5, -6, "total"),
        ]),
        (14, "Cash flow statement", [
            ("Net income (top line)", "={c}12", -7.5, -6, "formula"),
            ("Add back depreciation", "={c}9", 10, 10, "formula"),
            ("Cash from operations", "={c}15+{c}16", 2.5, 4, "total"),
            ("Net change in cash", "={c}17", 2.5, 4, "total"),
        ]),
        (20, "Balance sheet", [
            ("Cash", "={c}18", 2.5, 4, "formula"),
            ("Net PP&E", "=-{c}9", -10, -10, "formula"),
            ("Total assets", "={c}21+{c}22", -7.5, -6, "total"),
            ("Retained earnings", "={c}12", -7.5, -6, "formula"),
            ("Total liabilities and equity", "={c}24", -7.5, -6, "total"),
            ("Balance check (assets minus liabilities and equity)", "={c}23-{c}25", 0, 0, "formula"),
        ]),
    ]
    for hdr_row, text, rows in blocks:
        section(ws, hdr_row, text, {"D": "25% tax rate", "E": "40% tax rate", "G": "Formula (25%)"}, last_col=7)
        r = hdr_row + 1
        for label, fml, want25, want40, kind in rows:
            put(ws, f"B{r}", label, bold=(kind == "total"), indent=0 if kind == "total" else 1)
            cell(ws, f"D{r}", fml.format(c="D"), want25, kind=kind, fmt=FMT_DEC)
            cell(ws, f"E{r}", fml.format(c="E"), want40, kind=kind, fmt=FMT_DEC)
            formula_text(ws, f"G{r}", fml.format(c="D"))
            r += 1
    from openpyxl.styles import PatternFill
    for col in "BDE":
        ws[f"{col}26"].fill = PatternFill("solid", fgColor=FILL_SUBTLE)
    note(ws, "B28", "The 2.5 of cash is the tax shield: 10 of depreciation times the 25 percent rate.")

    # the typed article values above against the drafting script
    assert (w25["ni"], w25["cfo"], w25["cash"], w25["assets"], w25["check"]) == (-7.5, 2.5, 2.5, -7.5, 0)
    assert (w40["ni"], w40["cfo"], w40["cash"], w40["assets"], w40["check"]) == (-6, 4, 4, -6, 0)
    return ws


# ---------------------------------------------------------------- sheet 3: the ten scenarios
# Each block: title, inputs [(label, value)], rows [(label, formula or None for "no change", article value, kind)].
# Formulas use {t} for the shared tax rate cell and {i0}, {i1} for the block's input cells, {r0}.. for the
# block's own rows in order. The article's tables are reproduced row for row.
SCENARIOS = [
    ("Inventory goes up by $10, paid in cash", FMT_INT,
     [("Inventory bought, paid in cash", 10)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: change in inventory (asset up, cash out)", "=-{i0}", -10, "formula"),
      ("Cash flow statement: cash from operations", "={r1}", -10, "total"),
      ("Cash flow statement: net change in cash", "={r2}", -10, "total"),
      ("Balance sheet: cash", "={r3}", -10, "formula"),
      ("Balance sheet: inventory", "={i0}", 10, "formula"),
      ("Balance sheet: total assets", "={r4}+{r5}", 0, "total"),
      ("Balance sheet: total liabilities and equity", "=0", 0, "total"),
      ("Balance check", "={r6}-{r7}", 0, "formula")]),
    ("Inventory goes up by $10, bought on credit", FMT_INT,
     [("Inventory bought on credit (supplier not yet paid)", 10)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: change in inventory (asset up, cash out)", "=-{i0}", -10, "formula"),
      ("Cash flow statement: change in accounts payable (liability up, cash in)", "={i0}", 10, "formula"),
      ("Cash flow statement: cash from operations", "={r1}+{r2}", 0, "total"),
      ("Cash flow statement: net change in cash", "={r3}", 0, "total"),
      ("Balance sheet: inventory", "={i0}", 10, "formula"),
      ("Balance sheet: total assets", "={r4}+{r5}", 10, "total"),
      ("Balance sheet: accounts payable", "={i0}", 10, "formula"),
      ("Balance sheet: total liabilities and equity", "={r7}", 10, "total"),
      ("Balance check", "={r6}-{r8}", 0, "formula")]),
    ("$100 of capex, paid in cash", FMT_INT,
     [("Capital expenditure, paid in cash", 100)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: cash from operations", "=0", 0, "formula"),
      ("Cash flow statement: capital expenditures (investing)", "=-{i0}", -100, "formula"),
      ("Cash flow statement: net change in cash", "={r1}+{r2}", -100, "total"),
      ("Balance sheet: cash", "={r3}", -100, "formula"),
      ("Balance sheet: net PP&E", "={i0}", 100, "formula"),
      ("Balance sheet: total assets", "={r4}+{r5}", 0, "total"),
      ("Balance sheet: total liabilities and equity", "=0", 0, "total"),
      ("Balance check", "={r6}-{r7}", 0, "formula")]),
    ("Raise $100 of debt", FMT_INT,
     [("Debt raised", 100)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: cash from operations", "=0", 0, "formula"),
      ("Cash flow statement: debt issued (financing)", "={i0}", 100, "formula"),
      ("Cash flow statement: net change in cash", "={r1}+{r2}", 100, "total"),
      ("Balance sheet: cash", "={r3}", 100, "formula"),
      ("Balance sheet: total assets", "={r4}", 100, "total"),
      ("Balance sheet: long-term debt", "={i0}", 100, "formula"),
      ("Balance sheet: total liabilities and equity", "={r6}", 100, "total"),
      ("Balance check", "={r5}-{r7}", 0, "formula")]),
    ("Buy back $50 of stock", FMT_INT,
     [("Shares repurchased for cash", 50)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: cash from operations", "=0", 0, "formula"),
      ("Cash flow statement: share repurchases (financing)", "=-{i0}", -50, "formula"),
      ("Cash flow statement: net change in cash", "={r1}+{r2}", -50, "total"),
      ("Balance sheet: cash", "={r3}", -50, "formula"),
      ("Balance sheet: total assets", "={r4}", -50, "total"),
      ("Balance sheet: treasury stock (reduces equity)", "=-{i0}", -50, "formula"),
      ("Balance sheet: total liabilities and equity", "={r6}", -50, "total"),
      ("Balance check", "={r5}-{r7}", 0, "formula")]),
    ("$20 of deferred revenue collected", FMT_INT,
     [("Cash collected before the service is delivered", 20)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: change in deferred revenue (liability up, cash in)", "={i0}", 20, "formula"),
      ("Cash flow statement: cash from operations", "={r1}", 20, "total"),
      ("Cash flow statement: net change in cash", "={r2}", 20, "total"),
      ("Balance sheet: cash", "={r3}", 20, "formula"),
      ("Balance sheet: total assets", "={r4}", 20, "total"),
      ("Balance sheet: deferred revenue", "={i0}", 20, "formula"),
      ("Balance sheet: total liabilities and equity", "={r6}", 20, "total"),
      ("Balance check", "={r5}-{r7}", 0, "formula")]),
    ("A $50 write-down of PP&E (treated as tax deductible)", FMT_DEC,
     [("Write-down of PP&E", 50), ("Tax rate", "=$D$3")],
     [("Income statement: write-down expense", "={i0}", 50, "formula"),
      ("Income statement: pre-tax income", "=-{r0}", -50, "formula"),
      ("Income statement: taxes", "={r1}*{t}", -12.5, "formula"),
      ("Income statement: net income", "={r1}-{r2}", -37.5, "total"),
      ("Cash flow statement: net income", "={r3}", -37.5, "formula"),
      ("Cash flow statement: add back write-down (non-cash)", "={r0}", 50, "formula"),
      ("Cash flow statement: cash from operations", "={r4}+{r5}", 12.5, "total"),
      ("Cash flow statement: net change in cash", "={r6}", 12.5, "total"),
      ("Balance sheet: cash", "={r7}", 12.5, "formula"),
      ("Balance sheet: net PP&E", "=-{i0}", -50, "formula"),
      ("Balance sheet: total assets", "={r8}+{r9}", -37.5, "total"),
      ("Balance sheet: retained earnings", "={r3}", -37.5, "formula"),
      ("Balance sheet: total liabilities and equity", "={r11}", -37.5, "total"),
      ("Balance check", "={r10}-{r12}", 0, "formula")]),
    ("A $10 dividend", FMT_INT,
     [("Dividend paid in cash", 10)],
     [("Income statement", None, None, "formula"),
      ("Cash flow statement: cash from operations", "=0", 0, "formula"),
      ("Cash flow statement: dividends paid (financing)", "=-{i0}", -10, "formula"),
      ("Cash flow statement: net change in cash", "={r1}+{r2}", -10, "total"),
      ("Balance sheet: cash", "={r3}", -10, "formula"),
      ("Balance sheet: total assets", "={r4}", -10, "total"),
      ("Balance sheet: retained earnings", "=-{i0}", -10, "formula"),
      ("Balance sheet: total liabilities and equity", "={r6}", -10, "total"),
      ("Balance check", "={r5}-{r7}", 0, "formula")]),
    ("Accrued expenses go up by $10", FMT_DEC,
     [("Expense recognized but not yet paid", 10), ("Tax rate", "=$D$3")],
     [("Income statement: operating expense", "={i0}", 10, "formula"),
      ("Income statement: pre-tax income", "=-{r0}", -10, "formula"),
      ("Income statement: taxes", "={r1}*{t}", -2.5, "formula"),
      ("Income statement: net income", "={r1}-{r2}", -7.5, "total"),
      ("Cash flow statement: net income", "={r3}", -7.5, "formula"),
      ("Cash flow statement: change in accrued expenses (liability up, cash in)", "={i0}", 10, "formula"),
      ("Cash flow statement: cash from operations", "={r4}+{r5}", 2.5, "total"),
      ("Cash flow statement: net change in cash", "={r6}", 2.5, "total"),
      ("Balance sheet: cash", "={r7}", 2.5, "formula"),
      ("Balance sheet: total assets", "={r8}", 2.5, "total"),
      ("Balance sheet: accrued expenses", "={i0}", 10, "formula"),
      ("Balance sheet: retained earnings", "={r3}", -7.5, "formula"),
      ("Balance sheet: total liabilities and equity", "={r10}+{r11}", 2.5, "total"),
      ("Balance check", "={r9}-{r12}", 0, "formula")]),
    ("Sell equipment with a $30 book value for $40", FMT_DEC,
     [("Book value of the equipment", 30), ("Cash received for it", 40), ("Tax rate", "=$D$3")],
     [("Income statement: gain on sale", "={i1}-{i0}", 10, "formula"),
      ("Income statement: pre-tax income", "={r0}", 10, "formula"),
      ("Income statement: taxes", "=-{r1}*{t}", -2.5, "formula"),
      ("Income statement: net income", "={r1}+{r2}", 7.5, "total"),
      ("Cash flow statement: net income", "={r3}", 7.5, "formula"),
      ("Cash flow statement: subtract gain on sale (non-cash)", "=-{r0}", -10, "formula"),
      ("Cash flow statement: cash from operations", "={r4}+{r5}", -2.5, "total"),
      ("Cash flow statement: proceeds from sale of equipment (investing)", "={i1}", 40, "formula"),
      ("Cash flow statement: net change in cash", "={r6}+{r7}", 37.5, "total"),
      ("Balance sheet: cash", "={r8}", 37.5, "formula"),
      ("Balance sheet: net PP&E", "=-{i0}", -30, "formula"),
      ("Balance sheet: total assets", "={r9}+{r10}", 7.5, "total"),
      ("Balance sheet: retained earnings", "={r3}", 7.5, "formula"),
      ("Balance sheet: total liabilities and equity", "={r12}", 7.5, "total"),
      ("Balance check", "={r11}-{r13}", 0, "formula")]),
]

SCENARIO_RANGES = {}   # title -> (first_row, last_row) on the Scenarios sheet, for the screenshots


def build_scenarios(wb):
    ws = wb.create_sheet("Scenarios")
    set_widths(ws, {"A": 2, "B": 66, "C": 2, "D": 11, "E": 2, "F": 22})
    ws.sheet_view.showGridLines = False
    title(ws, "B1", "Ten more scenarios, each worked the same way")
    note(ws, "B2", "$. Every table is the change caused by the scenario. Parentheses mean the line went down. "
                   "The tax rate below feeds every block.")
    put(ws, "B3", "Tax rate", indent=1)
    put(ws, "D3", 0.25, kind="input", fmt=FMT_PCT0)
    tax_ref = "$D$3"

    from openpyxl.styles import PatternFill
    r = 5
    for name, fmt, inputs, rows in SCENARIOS:
        first = r
        section(ws, r, name, {"D": "Change", "F": "Formula"}, last_col=6)
        r += 1
        input_refs, block_tax = [], tax_ref
        for label, val in inputs:
            if isinstance(val, str):          # the shared tax rate, shown inside the block as a link
                put(ws, f"B{r}", f"{label} (from the input above)", italic=True, indent=1)
                put(ws, f"D{r}", val, kind="link", fmt=FMT_PCT0)
                block_tax = f"D{r}"
            else:
                put(ws, f"B{r}", f"Input: {label}", italic=True, indent=1)
                put(ws, f"D{r}", val, kind="input", fmt=fmt)
                input_refs.append(f"D{r}")
            r += 1
        row_refs = []
        for label, fml, want, kind in rows:
            put(ws, f"B{r}", label, bold=(kind == "total" or label == "Balance check"),
                indent=0 if kind == "total" else 1)
            if fml is None:
                no_change(ws, f"D{r}")
            else:
                subs = {"t": block_tax}
                subs.update({f"i{k}": ref for k, ref in enumerate(input_refs)})
                subs.update({f"r{k}": ref for k, ref in enumerate(row_refs)})
                f = fml.format(**subs)
                cell(ws, f"D{r}", f, want, kind=kind, fmt=fmt)
                formula_text(ws, f"F{r}", f)
            if label == "Balance check":
                for col in "BD":
                    ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=FILL_SUBTLE)
            row_refs.append(f"D{r}")
            r += 1
        SCENARIO_RANGES[name] = (first, r - 1)
        r += 1   # one blank row between blocks
    return ws


# ---------------------------------------------------------------- build and check
def build(out=OUT):
    EXPECTED.clear()
    SCENARIO_RANGES.clear()
    numbers = load_numbers()
    wb = new_workbook()
    wb.properties.creator = "Wall Street Guide"
    wb.properties.lastModifiedBy = "Wall Street Guide"
    build_lantern(wb, numbers)
    build_depreciation(wb, numbers)
    build_scenarios(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def check_against_article(xlsx=OUT, verbose=True):
    """Recalculate with LibreOffice and compare every cell written with an article value."""
    if not EXPECTED:
        build(xlsx)
    vals = recalc_values(xlsx)
    mismatches, n = [], 0
    for sheet, cells in EXPECTED.items():
        for ref, want in cells.items():
            got = vals[sheet].get(ref)
            n += 1
            if not isinstance(got, (int, float)) or round(float(got), 6) != round(float(want), 6):
                mismatches.append((sheet, ref, want, got))
    if verbose:
        print(f"compared {n} cells against the article's tables; mismatches: {len(mismatches)}")
        for m in mismatches:
            print("  MISMATCH", m)
    return mismatches, vals


if __name__ == "__main__":
    path = build()
    print("built", path)
    mm, vals = check_against_article(path)
    assert not mm, "workbook disagrees with the article"
    lb = vals["Lantern Bikes"]
    assert lb[f"D{R['check']}"] == 0 and lb[f"E{R['check']}"] == 0, "Lantern Bikes does not balance"
    print("balance check: Year 1", lb[f"D{R['check']}"], "Year 2", lb[f"E{R['check']}"])
