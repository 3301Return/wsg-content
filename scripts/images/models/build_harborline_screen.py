#!/usr/bin/env python3
"""Build scripts/images/models/harborline-screen.xlsx: the Harborline investment screen from
"Growth Equity Interview Guide for Undergrads" (October 5, 2026), with live formulas.

Harborline is fictional (dispatch and billing software for freight brokers). Three sheets:

  Metrics   section 7: the ARR bridge, growth and retention, the cost inputs and the unit
            economics (CAC, CAC payback, LTV, LTV to CAC, magic number, burn multiple,
            free cash flow margin, Rule of 40)
  Deal      section 9: the round (new money, post and pre-money, ownership, price per share,
            new shares) and the cap table before and after
  Returns   section 11: three exits five years out, the payout to the preferred three ways
            (as common, 1x non-participating, 1x participating), MOIC and IRR, what the
            common holders keep, the $500.0 million post-money price check, and the payout
            formulas written out as text for the close-up

Inputs come from scripts/images/models/numbers_growth-equity-interview-guide-undergrads.py
(the plain-Python source of every figure in the article), run silently and read here, so the
workbook can never drift from it. Blue on yellow = typed input, black = formula, green = a value
pulled from another sheet.

Run from the repo root:  python3 scripts/images/models/build_harborline_screen.py
Idempotent: rewrites the workbook every time. Workbook author: Wall Street Guide.
"""
import contextlib
import io
import runpy
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts" / "images"))
from wsg_viz import XL_FONT, XL_SIZE, header_bar, new_workbook, put, set_widths  # noqa: E402

OUT_XLSX = ROOT / "scripts" / "images" / "models" / "harborline-screen.xlsx"
NUMBERS = ROOT / "scripts" / "images" / "models" / "numbers_growth-equity-interview-guide-undergrads.py"

FMT_M = '"$"#,##0.0;("$"#,##0.0);"$"0.0'          # $ millions, one decimal, negatives in parentheses
FMT_D0 = '"$"#,##0;("$"#,##0);"$"0'                # whole dollars
FMT_D2 = '"$"#,##0.00'
FMT_P1 = '0.0%;(0.0%);0.0%'
FMT_P2 = '0.00%;(0.00%);0.00%'
FMT_X1 = '0.0"x"'
FMT_X2 = '0.00"x"'
FMT_SH = '#,##0;(#,##0);0'
FMT_MONTHS = '0.0" months"'
FMT_R40 = '0'


def load_numbers():
    """Run the article's numbers script without its printout and return its globals."""
    with contextlib.redirect_stdout(io.StringIO()):
        return runpy.run_path(str(NUMBERS))


def col_headers(ws, row, items, wrap=True, height=None, left=("B",)):
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin", color="000000")
    for col, text in items.items():
        c = ws[f"{col}{row}"]
        c.value = text
        c.font = Font(name=XL_FONT, size=XL_SIZE, bold=True)
        c.alignment = Alignment(horizontal="left" if col in left else "right", vertical="bottom", wrap_text=wrap)
        c.border = Border(bottom=thin)
    if height:
        ws.row_dimensions[row].height = height


def note(ws, ref, text):
    return put(ws, ref, text, italic=True, color="52514E")


def formula_text(ws, ref, text):
    c = put(ws, ref, text, italic=True, color="52514E")
    c.data_type = "s"
    return c


def build(out_path=OUT_XLSX):
    G = load_numbers()
    wb = new_workbook()
    wb.properties.creator = "Wall Street Guide"
    wb.properties.lastModifiedBy = "Wall Street Guide"
    wb.properties.title = "Harborline growth equity screen"
    M, D, R = {}, {}, {}

    # ================================================================ Metrics
    ws = wb.active
    ws.title = "Metrics"
    set_widths(ws, {"A": 2, "B": 44, "C": 2, "D": 13, "E": 2, "F": 50})
    ws.sheet_view.showGridLines = False
    put(ws, "B2", "Harborline (fictional): ARR bridge, retention and unit economics", bold=True, size=12)
    note(ws, "B3", "$ millions unless a line says otherwise. Blue on yellow = typed input, black = formula.")

    r = 5
    M["bridge_header"] = r
    header_bar(ws, r, 2, 6, "ARR BRIDGE, LAST TWELVE MONTHS"); r += 1
    col_headers(ws, r, {"B": "Line", "D": "$ millions", "F": "How it is calculated"}, left=("B", "F")); r += 1
    M["begin"] = r; put(ws, f"B{r}", "Beginning ARR (start of year)"); put(ws, f"D{r}", G["BEGIN_ARR"], kind="input", fmt=FMT_M); note(ws, f"F{r}", "given"); r += 1
    M["churn"] = r; put(ws, f"B{r}", "Minus churned ARR"); put(ws, f"D{r}", -G["CHURNED_ARR"], kind="input", fmt=FMT_M); note(ws, f"F{r}", "customers that left or downgraded"); r += 1
    M["exp"] = r; put(ws, f"B{r}", "Plus expansion ARR"); put(ws, f"D{r}", G["EXPANSION_ARR"], kind="input", fmt=FMT_M); note(ws, f"F{r}", "existing customers buying more"); r += 1
    M["new"] = r; put(ws, f"B{r}", "Plus new logo ARR"); put(ws, f"D{r}", G["NEW_LOGO_ARR"], kind="input", fmt=FMT_M); note(ws, f"F{r}", "brand-new customers"); r += 1
    M["end"] = r; put(ws, f"B{r}", "Ending ARR (today)", kind="total"); put(ws, f"D{r}", f"=SUM(D{M['begin']}:D{M['new']})", kind="total", fmt=FMT_M); note(ws, f"F{r}", "40.0 minus 4.0 plus 6.0 plus 18.0"); r += 1
    M["net_new"] = r; put(ws, f"B{r}", "Net new ARR"); put(ws, f"D{r}", f"=D{M['end']}-D{M['begin']}", kind="formula", fmt=FMT_M); note(ws, f"F{r}", "60.0 minus 40.0"); r += 1
    M["gross_new"] = r; put(ws, f"B{r}", "Gross new ARR (new plus expansion)"); put(ws, f"D{r}", f"=D{M['new']}+D{M['exp']}", kind="formula", fmt=FMT_M); note(ws, f"F{r}", "18.0 plus 6.0"); r += 2

    M["ret_header"] = r
    header_bar(ws, r, 2, 6, "GROWTH AND RETENTION FROM THE BRIDGE"); r += 1
    col_headers(ws, r, {"B": "Metric", "D": "Value", "F": "Arithmetic"}, left=("B", "F")); r += 1
    M["growth"] = r; put(ws, f"B{r}", "ARR growth"); put(ws, f"D{r}", f"=D{M['end']}/D{M['begin']}-1", kind="formula", fmt=FMT_P1); note(ws, f"F{r}", "60.0 divided by 40.0, minus 1"); r += 1
    M["grr"] = r; put(ws, f"B{r}", "Gross revenue retention"); put(ws, f"D{r}", f"=(D{M['begin']}+D{M['churn']})/D{M['begin']}", kind="formula", fmt=FMT_P1); note(ws, f"F{r}", "(40.0 minus 4.0) divided by 40.0"); r += 1
    M["nrr"] = r; put(ws, f"B{r}", "Net revenue retention"); put(ws, f"D{r}", f"=(D{M['begin']}+D{M['churn']}+D{M['exp']})/D{M['begin']}", kind="formula", fmt=FMT_P1); note(ws, f"F{r}", "(40.0 minus 4.0 plus 6.0) divided by 40.0"); r += 1
    M["bridge_end"] = r - 1
    r += 1

    M["ue_header"] = r
    header_bar(ws, r, 2, 6, "COST INPUTS"); r += 1
    col_headers(ws, r, {"B": "Input", "D": "Value", "F": "Note"}, left=("B", "F")); r += 1
    M["rev"] = r; put(ws, f"B{r}", "Revenue this year"); put(ws, f"D{r}", G["REVENUE"], kind="input", fmt=FMT_M); note(ws, f"F{r}", "ARR ramped from 40.0 to 60.0 during the year"); r += 1
    M["gm"] = r; put(ws, f"B{r}", "Gross margin"); put(ws, f"D{r}", G["GROSS_MARGIN"], kind="input", fmt=FMT_P1); note(ws, f"F{r}", "after hosting, support and implementation"); r += 1
    M["gp"] = r; put(ws, f"B{r}", "Gross profit"); put(ws, f"D{r}", f"=D{M['rev']}*D{M['gm']}", kind="formula", fmt=FMT_M); note(ws, f"F{r}", "50.0 times 76%"); r += 1
    M["sm"] = r; put(ws, f"B{r}", "Sales and marketing spend"); put(ws, f"D{r}", G["SALES_MARKETING"], kind="input", fmt=FMT_M); r += 1
    M["burn"] = r; put(ws, f"B{r}", "Net burn (cash out minus cash in)"); put(ws, f"D{r}", G["NET_BURN"], kind="input", fmt=FMT_M); r += 1
    M["arr_cust"] = r; put(ws, f"B{r}", "ARR per customer ($)"); put(ws, f"D{r}", G["ARR_PER_CUSTOMER"] * 1e6, kind="input", fmt=FMT_D0); r += 1
    M["cust"] = r; put(ws, f"B{r}", "Customers today"); put(ws, f"D{r}", f"=D{M['end']}*1000000/D{M['arr_cust']}", kind="formula", fmt=FMT_SH); note(ws, f"F{r}", "ending ARR divided by ARR per customer"); r += 1
    M["new_cust"] = r; put(ws, f"B{r}", "New customers signed this year"); put(ws, f"D{r}", f"=D{M['new']}*1000000/D{M['arr_cust']}", kind="formula", fmt=FMT_SH); note(ws, f"F{r}", "new logo ARR divided by ARR per customer"); r += 2

    M["uem_header"] = r
    header_bar(ws, r, 2, 6, "UNIT ECONOMICS"); r += 1
    col_headers(ws, r, {"B": "Metric", "D": "Value", "F": "Arithmetic"}, left=("B", "F")); r += 1
    M["cac"] = r; put(ws, f"B{r}", "CAC (cost to win one customer, $)"); put(ws, f"D{r}", f"=D{M['sm']}*1000000/D{M['new_cust']}", kind="formula", fmt=FMT_D0); note(ws, f"F{r}", "$20.0M divided by 300 new customers"); r += 1
    M["payback"] = r; put(ws, f"B{r}", "CAC payback"); put(ws, f"D{r}", f"=D{M['sm']}/(D{M['gross_new']}*D{M['gm']})*12", kind="formula", fmt=FMT_MONTHS); note(ws, f"F{r}", "$20.0M divided by ($24.0M times 76%), times 12"); r += 1
    M["ltv"] = r; put(ws, f"B{r}", "LTV (lifetime gross profit per customer, $)"); put(ws, f"D{r}", f"=D{M['arr_cust']}*D{M['gm']}/(-D{M['churn']}/D{M['begin']})", kind="formula", fmt=FMT_D0); note(ws, f"F{r}", "$60,000 times 76%, divided by 10% churn"); r += 1
    M["ltv_cac"] = r; put(ws, f"B{r}", "LTV to CAC"); put(ws, f"D{r}", f"=D{M['ltv']}/D{M['cac']}", kind="formula", fmt=FMT_X1); note(ws, f"F{r}", "$456,000 divided by $66,667"); r += 1
    M["magic"] = r; put(ws, f"B{r}", "Magic number"); put(ws, f"D{r}", f"=D{M['net_new']}/D{M['sm']}", kind="formula", fmt=FMT_X2); note(ws, f"F{r}", "$20.0M net new ARR divided by $20.0M S&M"); r += 1
    M["burn_mult"] = r; put(ws, f"B{r}", "Burn multiple"); put(ws, f"D{r}", f"=D{M['burn']}/D{M['net_new']}", kind="formula", fmt=FMT_X2); note(ws, f"F{r}", "$5.0M net burn divided by $20.0M net new ARR"); r += 1
    M["fcf"] = r; put(ws, f"B{r}", "Free cash flow margin"); put(ws, f"D{r}", f"=-D{M['burn']}/D{M['rev']}", kind="formula", fmt=FMT_P1); note(ws, f"F{r}", "negative $5.0M divided by $50.0M"); r += 1
    M["r40"] = r; put(ws, f"B{r}", "Rule of 40 score"); put(ws, f"D{r}", f"=(D{M['growth']}+D{M['fcf']})*100", kind="formula", fmt=FMT_R40); note(ws, f"F{r}", "50.0% growth plus (10.0%) margin"); r += 1
    M["ue_end"] = r - 1

    # ================================================================ Deal
    ds = wb.create_sheet("Deal")
    set_widths(ds, {"A": 2, "B": 40, "C": 2, "D": 15, "E": 11, "F": 15, "G": 11})
    ds.sheet_view.showGridLines = False
    put(ds, "B2", "Harborline Series C: the round and the cap table", bold=True, size=12)
    note(ds, "B3", "$ millions unless a line says otherwise. Green = pulled from the Metrics sheet.")
    r = 5
    D["terms_header"] = r
    header_bar(ds, r, 2, 7, "THE ROUND"); r += 1
    D["inv"] = r; put(ds, f"B{r}", "New money in"); put(ds, f"D{r}", G["INVESTMENT"], kind="input", fmt=FMT_M); r += 1
    D["post"] = r; put(ds, f"B{r}", "Post-money valuation"); put(ds, f"D{r}", G["POST_MONEY"], kind="input", fmt=FMT_M); r += 1
    D["pre"] = r; put(ds, f"B{r}", "Pre-money valuation"); put(ds, f"D{r}", f"=D{D['post']}-D{D['inv']}", kind="formula", fmt=FMT_M); r += 1
    D["own"] = r; put(ds, f"B{r}", "Ownership bought"); put(ds, f"D{r}", f"=D{D['inv']}/D{D['post']}", kind="formula", fmt=FMT_P2); r += 1
    D["arr"] = r; put(ds, f"B{r}", "Ending ARR (Metrics sheet)"); put(ds, f"D{r}", f"=Metrics!D{M['end']}", kind="link", fmt=FMT_M); r += 1
    D["ev_arr"] = r; put(ds, f"B{r}", "Entry EV to ARR"); put(ds, f"D{r}", f"=D{D['post']}/D{D['arr']}", kind="formula", fmt=FMT_X1); r += 1
    D["pre_sh"] = r; put(ds, f"B{r}", "Shares before the round"); r += 1   # value written below, once the cap table total row is known
    D["pps"] = r; put(ds, f"B{r}", "Price per share ($)"); put(ds, f"D{r}", f"=D{D['pre']}*1000000/D{D['pre_sh']}", kind="formula", fmt=FMT_D2); r += 1
    D["new_sh"] = r; put(ds, f"B{r}", "New shares issued"); put(ds, f"D{r}", f"=D{D['inv']}*1000000/D{D['pps']}", kind="formula", fmt=FMT_SH); r += 1
    D["post_sh"] = r; put(ds, f"B{r}", "Shares after the round", kind="total"); put(ds, f"D{r}", f"=D{D['pre_sh']}+D{D['new_sh']}", kind="total", fmt=FMT_SH); r += 2

    D["cap_header"] = r
    header_bar(ds, r, 2, 7, "CAP TABLE BEFORE AND AFTER (fully diluted, option pool included)"); r += 1
    col_headers(ds, r, {"B": "Holder", "D": "Shares before", "E": "% before", "F": "Shares after", "G": "% after"}); r += 1
    D["cap_first"] = r
    holders = [("Founders", 30_000_000), ("Earlier venture investors", 15_000_000), ("Employee option pool", 5_000_000)]
    for name, sh in holders:
        put(ds, f"B{r}", name)
        put(ds, f"D{r}", sh, kind="input", fmt=FMT_SH)
        put(ds, f"F{r}", f"=D{r}", kind="formula", fmt=FMT_SH)
        r += 1
    D["cap_inv"] = r
    put(ds, f"B{r}", "New growth investor (Series C preferred)")
    put(ds, f"D{r}", 0, kind="input", fmt=FMT_SH)
    put(ds, f"F{r}", f"=D{D['new_sh']}", kind="formula", fmt=FMT_SH)
    r += 1
    D["cap_total"] = r
    put(ds, f"B{r}", "Total", kind="total")
    put(ds, f"D{r}", f"=SUM(D{D['cap_first']}:D{D['cap_inv']})", kind="total", fmt=FMT_SH)
    put(ds, f"F{r}", f"=SUM(F{D['cap_first']}:F{D['cap_inv']})", kind="total", fmt=FMT_SH)
    for rr in range(D["cap_first"], r + 1):
        kind = "total" if rr == r else "formula"
        put(ds, f"E{rr}", f"=D{rr}/D${D['cap_total']}", kind=kind, fmt='0.0%')
        put(ds, f"G{rr}", f"=F{rr}/F${D['cap_total']}", kind=kind, fmt='0.00%')
    r += 1
    D["cap_check"] = r
    note(ds, f"B{r}", "Founders after, the shortcut: 60.0% times pre over post")
    put(ds, f"G{r}", f"=E{D['cap_first']}*D{D['pre']}/D{D['post']}", kind="formula", fmt='0.00%', italic=True)
    D["cap_end"] = r
    # shares before the round = the cap table's total
    put(ds, f"D{D['pre_sh']}", f"=D{D['cap_total']}", kind="formula", fmt=FMT_SH)

    # ================================================================ Returns
    rs = wb.create_sheet("Returns")
    set_widths(rs, {"A": 2, "B": 30, "C": 2, "D": 14, "E": 14, "F": 14, "G": 14, "H": 18, "I": 14, "J": 2})
    rs.sheet_view.showGridLines = False
    put(rs, "B2", "Harborline screen: three exits, the preference, MOIC and IRR", bold=True, size=12)
    note(rs, "B3", "$ millions. Five-year hold, one check in and one check out; equity value equals EV at exit (no debt or cash).")
    r = 5
    R["in_header"] = r
    header_bar(rs, r, 2, 9, "DEAL INPUTS"); r += 1
    R["inv"] = r; put(rs, f"B{r}", "Investment (Deal sheet)"); put(rs, f"D{r}", f"=Deal!D{D['inv']}", kind="link", fmt=FMT_M); r += 1
    R["own"] = r; put(rs, f"B{r}", "Ownership (Deal sheet)"); put(rs, f"D{r}", f"=Deal!D{D['own']}", kind="link", fmt=FMT_P2); r += 1
    R["arr0"] = r; put(rs, f"B{r}", "ARR today (Metrics sheet)"); put(rs, f"D{r}", f"=Metrics!D{M['end']}", kind="link", fmt=FMT_M); r += 1
    R["pref_x"] = r; put(rs, f"B{r}", "Liquidation preference (x invested)"); put(rs, f"D{r}", G["PREF_MULTIPLE"], kind="input", fmt=FMT_X1); r += 1
    R["pref"] = r; put(rs, f"B{r}", "Preference amount"); put(rs, f"D{r}", f"=D{R['inv']}*D{R['pref_x']}", kind="formula", fmt=FMT_M); r += 1
    R["hold"] = r; put(rs, f"B{r}", "Hold (years)"); put(rs, f"D{r}", G["HOLD_YEARS"], kind="input", fmt='0'); r += 2

    names = [s[0] for s in G["scenarios"]]
    R["sc_header"] = r
    header_bar(rs, r, 2, 9, "STEP 3: THREE EXITS, FIVE YEARS OUT"); r += 1
    col_headers(rs, r, {"B": "Scenario", "D": "ARR in year 5", "E": "ARR CAGR from today", "F": "Exit EV to ARR",
                        "G": "Equity value at exit", "H": "Our stake as common"}, height=27); r += 1
    R["sc_first"] = r
    for name, arr5, mult5 in G["scenarios"]:
        put(rs, f"B{r}", name)
        put(rs, f"D{r}", arr5, kind="input", fmt=FMT_M)
        put(rs, f"E{r}", f"=(D{r}/$D${R['arr0']})^(1/$D${R['hold']})-1", kind="formula", fmt=FMT_P1)
        put(rs, f"F{r}", mult5, kind="input", fmt=FMT_X1)
        put(rs, f"G{r}", f"=D{r}*F{r}", kind="formula", fmt=FMT_M)
        put(rs, f"H{r}", f"=G{r}*$D${R['own']}", kind="formula", fmt=FMT_M)
        r += 1
    R["sc_end"] = r - 1
    r += 1

    R["po_header"] = r
    header_bar(rs, r, 2, 9, "STEP 4: WHAT THE PREFERRED GETS"); r += 1
    col_headers(rs, r, {"B": "Scenario", "D": "Equity value", "E": "As common (no preference)", "F": "1x preference amount",
                        "G": "1x non-participating payout", "H": "Which the investor picks", "I": "1x participating payout"}, height=40); r += 1
    R["po_first"] = r
    for i, name in enumerate(names):
        sr = R["sc_first"] + i
        put(rs, f"B{r}", name)
        put(rs, f"D{r}", f"=G{sr}", kind="formula", fmt=FMT_M)
        put(rs, f"E{r}", f"=D{r}*$D${R['own']}", kind="formula", fmt=FMT_M)
        put(rs, f"F{r}", f"=$D${R['pref']}", kind="formula", fmt=FMT_M)
        put(rs, f"G{r}", f"=MIN(D{r},MAX(F{r},E{r}))", kind="formula", fmt=FMT_M)
        put(rs, f"H{r}", f'=IF(F{r}>E{r},"preference","convert to common")', kind="formula", align="right")
        put(rs, f"I{r}", f"=MIN(F{r},D{r})+MAX(D{r}-F{r},0)*$D${R['own']}", kind="formula", fmt=FMT_M)
        r += 1
    R["po_end"] = r - 1
    r += 1

    R["mo_header"] = r
    header_bar(rs, r, 2, 9, "STEP 5: MOIC AND IRR ON THE $100.0 MILLION"); r += 1
    col_headers(rs, r, {"B": "Scenario", "D": "MOIC as common", "E": "IRR as common", "F": "MOIC, 1x non-part.",
                        "G": "IRR, 1x non-part.", "H": "MOIC, 1x participating", "I": "IRR, 1x participating"}, height=27); r += 1
    R["mo_first"] = r
    for i, name in enumerate(names):
        pr = R["po_first"] + i
        put(rs, f"B{r}", name)
        for mcol, icol, src in (("D", "E", "E"), ("F", "G", "G"), ("H", "I", "I")):
            put(rs, f"{mcol}{r}", f"={src}{pr}/$D${R['inv']}", kind="formula", fmt=FMT_X2)
            put(rs, f"{icol}{r}", f"={mcol}{r}^(1/$D${R['hold']})-1", kind="formula", fmt=FMT_P1)
        r += 1
    R["mo_end"] = r - 1
    r += 1

    R["wg_header"] = r
    header_bar(rs, r, 2, 9, "WHO GETS WHAT (1x non-participating)"); r += 1
    col_headers(rs, r, {"B": "Scenario", "D": "Equity value", "E": "To preferred", "F": "To founders, VCs and employees",
                        "G": "Their share"}, height=40); r += 1
    R["wg_first"] = r
    for i, name in enumerate(names):
        pr = R["po_first"] + i
        put(rs, f"B{r}", name)
        put(rs, f"D{r}", f"=D{pr}", kind="formula", fmt=FMT_M)
        put(rs, f"E{r}", f"=G{pr}", kind="formula", fmt=FMT_M)
        put(rs, f"F{r}", f"=D{r}-E{r}", kind="formula", fmt=FMT_M)
        put(rs, f"G{r}", f"=F{r}/D{r}", kind="formula", fmt=FMT_P1)
        r += 1
    R["wg_end"] = r - 1
    r += 1

    R["pc_header"] = r
    header_bar(rs, r, 2, 9, "STEP 6: THE PRICE CHECK, SAME CHECK AT A $500.0 MILLION POST"); r += 1
    col_headers(rs, r, {"B": "Item", "D": "At $600.0M post", "E": "At $500.0M post"}, height=27); r += 1
    R["pc_post"] = r; put(rs, f"B{r}", "Post-money valuation"); put(rs, f"D{r}", f"=Deal!D{D['post']}", kind="link", fmt=FMT_M); put(rs, f"E{r}", 500.0, kind="input", fmt=FMT_M); r += 1
    R["pc_own"] = r; put(rs, f"B{r}", "Ownership")
    for c in "DE":
        put(rs, f"{c}{r}", f"=$D${R['inv']}/{c}{R['pc_post']}", kind="formula", fmt=FMT_P2)
    r += 1
    base_row = R["sc_first"] + names.index("Base")
    R["pc_pay"] = r; put(rs, f"B{r}", "Base case payout (converts)")
    for c in "DE":
        put(rs, f"{c}{r}", f"=MAX($D${R['pref']},$G${base_row}*{c}{R['pc_own']})", kind="formula", fmt=FMT_M)
    r += 1
    R["pc_moic"] = r; put(rs, f"B{r}", "Base case MOIC")
    for c in "DE":
        put(rs, f"{c}{r}", f"={c}{R['pc_pay']}/$D${R['inv']}", kind="formula", fmt=FMT_X2)
    r += 1
    R["pc_irr"] = r; put(rs, f"B{r}", "Base case IRR")
    for c in "DE":
        put(rs, f"{c}{r}", f"={c}{R['pc_moic']}^(1/$D${R['hold']})-1", kind="formula", fmt=FMT_P1)
    r += 1
    R["pc_end"] = r - 1
    r += 1

    # formula close-up: the Bear row of the payout block, each formula written out as text
    R["fx_header"] = r
    header_bar(rs, r, 2, 9, "THE PREFERENCE IN FORMULAS (Bear row of step 4)"); r += 1
    col_headers(rs, r, {"B": "Column", "D": "Bear value", "E": "Formula in the cell"}, wrap=False, left=("B", "E")); r += 1
    bear = R["po_first"] + names.index("Bear")
    R["fx_first"] = r
    for label, col in (("As common (no preference)", "E"), ("1x preference amount", "F"), ("1x non-participating payout", "G"),
                       ("Which the investor picks", "H"), ("1x participating payout", "I")):
        put(rs, f"B{r}", label)
        put(rs, f"D{r}", f"={col}{bear}", kind="formula", fmt=FMT_M if col != "H" else None, align="right")
        src = rs[f"{col}{bear}"].value
        formula_text(rs, f"E{r}", f" {col}{bear}:  {src}")
        r += 1
    note(rs, f"B{r}", f"D{R['own']} is the 16.67% ownership, D{R['pref']} the $100.0 preference, and D{bear} the $270.0 Bear equity value.")
    R["fx_end"] = r

    wb.calculation.fullCalcOnLoad = True
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, M, D, R, G


if __name__ == "__main__":
    path, M, D, R, _ = build()
    print(f"wrote {path}")
    print("Metrics rows:", M)
    print("Deal rows:", D)
    print("Returns rows:", R)
