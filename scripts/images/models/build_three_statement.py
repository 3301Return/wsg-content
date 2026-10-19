#!/usr/bin/env python3
"""Build scripts/images/models/three-statement-test.xlsx for the WSG article
"Three Statement Model: How to Build One for a Modeling Test" (October 1, 2026).

One sheet ("Model") laid out like a two hour modeling test, every projected cell a live
formula: drivers block (with the circularity toggle at the top), income statement, PP&E
and working capital schedules, balance sheet with the check row, cash flow statement,
debt and revolver block. Years 2025A, 2026E, 2027E, 2028E across columns D to G, labels
in column B, column I carries the 2026E formula as text so a screenshot can show how the
cells connect.

Two more sheets reproduce the article's section 6 error patterns:
  "Error constant"  2025A net PP&E typed as 412.0 instead of 400.0 -> check row 12.0 every year.
  "Error growing"   the PP&E roll never subtracts D&A -> check row 44.0, 91.5, 141.9.

Run from the repo root:  python3 scripts/images/models/build_three_statement.py
The numbers are the article's own inputs (section 9); the script asserts the recalculated
workbook against the article's tables after saving.
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts" / "images"))
from wsg_viz import *  # noqa: E402,F401

OUT = ROOT / "scripts" / "images" / "models" / "three-statement-test.xlsx"
YEARS = ["2025A", "2026E", "2027E", "2028E"]
COLS = ["D", "E", "F", "G"]           # one column per year
FCOL = "I"                            # formula text column (2026E formulas, as text)
FMT_NUM = '#,##0.0;(#,##0.0);0.0'     # zeros read 0.0, as in the article's tables (never a dash)
FMT_CHECK = FMT_NUM                    # the check row must read 0.0
FMT_DAYS_HIST = '0.0'
FMT_DAYS = '0'
FMT_TOGGLE = '0'
RED = "E34948"

# ---------------------------------------------------------------- row map
R = dict(
    title=1, units=2,
    drv_hdr=4, toggle=5, growth=6, gm=7, sga=8, da=9, capex=10, dso=11, dio=12, dpo=13,
    tax=14, rdebt=15, rcash=16, amort=17, payout=18, mincash=19,
    is_hdr=21, rev=22, cogs=23, gp=24, sgax=25, ebitda=26, dax=27, ebit=28, intexp=29,
    intinc=30, ebt=31, taxes=32, ni=33, div=34,
    ppe_hdr=36, ppe_beg=37, ppe_capex=38, ppe_da=39, ppe_end=40,
    wc_hdr=42, wc_ar=43, wc_inv=44, wc_ap=45, d_ar=46, d_inv=47, d_ap=48,
    bs_hdr=50, cash=51, ar=52, inv=53, ppe=54, ta=55, ap=56, rev_bs=57, tl=58, cs=59, re=60,
    tle=61, check=62,
    cf_hdr=64, cf_ni=65, cf_da=66, cf_ar=67, cf_inv=68, cf_ap=69, cfo=70, cf_capex=71, cfi=72,
    cf_amort=73, cf_draw=74, cf_repay=75, cf_div=76, cff=77, netchg=78, begcash=79, endcash=80,
    cash_pre=81,
    dr_hdr=83, dr_min=84, dr_pre=85, rv_beg=86, rv_draw=87, rv_repay=88, rv_end=89,
    tl_beg=90, tl_amort=91, tl_end=92, debt=93,
)

# Historical (2025A) inputs from the article's section 9 tables, $ millions.
HIST_IS = dict(rev=1000.0, cogs=-600.0, sgax=-250.0, dax=-40.0, intexp=-18.0, intinc=2.0,
               taxes=-18.8, div=-18.8)
HIST_BS = dict(cash=80.0, ar=125.0, inv=100.0, ppe=400.0, ap=50.0, rev_bs=0.0, tl=300.0,
               cs=100.0, re=255.0)
HIST_CAPEX = 50.0
# Projection drivers (2026E, 2027E, 2028E).
DRV = dict(
    growth=[0.10, 0.08, 0.06], gm=[0.40] * 3, sga=[0.25] * 3, da=[0.04] * 3,
    capex=[0.05, 0.12, 0.05], dso=[45, 45, 45], dio=[60, 60, 60], dpo=[30, 30, 30],
    tax=[0.20] * 3, rdebt=[0.06] * 3, rcash=[0.02] * 3, amort=[30.0] * 3,
    payout=[0.25] * 3, mincash=[50.0] * 3,
)
TOGGLE_DEFAULT = 0   # beginning balance interest, the article's default for a timed test


def projected_formulas(c, p, variant=None):
    """Formulas for projected column c (prior column p). Returns {row_key: formula}."""
    f = {
        "rev": f"={p}{R['rev']}*(1+{c}{R['growth']})",
        "cogs": f"=-{c}{R['rev']}*(1-{c}{R['gm']})",
        "gp": f"=SUM({c}{R['rev']}:{c}{R['cogs']})",
        "sgax": f"=-{c}{R['rev']}*{c}{R['sga']}",
        "ebitda": f"=SUM({c}{R['gp']}:{c}{R['sgax']})",
        "dax": f"=-{c}{R['rev']}*{c}{R['da']}",
        "ebit": f"=SUM({c}{R['ebitda']}:{c}{R['dax']})",
        "intexp": f"=-{c}{R['rdebt']}*IF($D${R['toggle']}=1,AVERAGE({p}{R['debt']},{c}{R['debt']}),{p}{R['debt']})",
        "intinc": f"={c}{R['rcash']}*{p}{R['cash']}",
        "ebt": f"=SUM({c}{R['ebit']}:{c}{R['intinc']})",
        "taxes": f"=-MAX(0,{c}{R['ebt']})*{c}{R['tax']}",
        "ni": f"=SUM({c}{R['ebt']}:{c}{R['taxes']})",
        "div": f"=-{c}{R['ni']}*{c}{R['payout']}",
        "ppe_beg": f"={p}{R['ppe_end']}",
        "ppe_capex": f"={c}{R['rev']}*{c}{R['capex']}",
        "ppe_da": f"={c}{R['dax']}",
        "ppe_end": f"=SUM({c}{R['ppe_beg']}:{c}{R['ppe_da']})",
        "wc_ar": f"={c}{R['rev']}*{c}{R['dso']}/365",
        "wc_inv": f"=-{c}{R['cogs']}*{c}{R['dio']}/365",
        "wc_ap": f"=-{c}{R['cogs']}*{c}{R['dpo']}/365",
        "d_ar": f"=-({c}{R['wc_ar']}-{p}{R['wc_ar']})",
        "d_inv": f"=-({c}{R['wc_inv']}-{p}{R['wc_inv']})",
        "d_ap": f"={c}{R['wc_ap']}-{p}{R['wc_ap']}",
        "cash": f"={c}{R['endcash']}",
        "ar": f"={c}{R['wc_ar']}",
        "inv": f"={c}{R['wc_inv']}",
        "ppe": f"={c}{R['ppe_end']}",
        "ta": f"=SUM({c}{R['cash']}:{c}{R['ppe']})",
        "ap": f"={c}{R['wc_ap']}",
        "rev_bs": f"={c}{R['rv_end']}",
        "tl": f"={c}{R['tl_end']}",
        "cs": f"={p}{R['cs']}",
        "re": f"={p}{R['re']}+{c}{R['ni']}+{c}{R['div']}",
        "tle": f"=SUM({c}{R['ap']}:{c}{R['re']})",
        "check": f"={c}{R['ta']}-{c}{R['tle']}",
        "cf_ni": f"={c}{R['ni']}",
        "cf_da": f"=-{c}{R['dax']}",
        "cf_ar": f"={c}{R['d_ar']}",
        "cf_inv": f"={c}{R['d_inv']}",
        "cf_ap": f"={c}{R['d_ap']}",
        "cfo": f"=SUM({c}{R['cf_ni']}:{c}{R['cf_ap']})",
        "cf_capex": f"=-{c}{R['ppe_capex']}",
        "cfi": f"=SUM({c}{R['cf_capex']}:{c}{R['cf_capex']})",
        "cf_amort": f"=-{c}{R['amort']}",
        "cf_draw": f"={c}{R['rv_draw']}",
        "cf_repay": f"=-{c}{R['rv_repay']}",
        "cf_div": f"={c}{R['div']}",
        "cff": f"=SUM({c}{R['cf_amort']}:{c}{R['cf_div']})",
        "netchg": f"={c}{R['cfo']}+{c}{R['cfi']}+{c}{R['cff']}",
        "begcash": f"={p}{R['cash']}",
        "endcash": f"={c}{R['begcash']}+{c}{R['netchg']}",
        "cash_pre": f"={c}{R['begcash']}+{c}{R['cfo']}+{c}{R['cfi']}+{c}{R['cf_amort']}+{c}{R['cf_div']}",
        "dr_min": f"={c}{R['mincash']}",
        "dr_pre": f"={c}{R['cash_pre']}",
        "rv_beg": f"={p}{R['rv_end']}",
        "rv_draw": f"=MAX(0,{c}{R['dr_min']}-{c}{R['dr_pre']})",
        "rv_repay": f"=MIN({c}{R['rv_beg']},MAX(0,{c}{R['dr_pre']}-{c}{R['dr_min']}))",
        "rv_end": f"={c}{R['rv_beg']}+{c}{R['rv_draw']}-{c}{R['rv_repay']}",
        "tl_beg": f"={p}{R['tl_end']}",
        "tl_amort": f"=-{c}{R['amort']}",
        "tl_end": f"=SUM({c}{R['tl_beg']}:{c}{R['tl_amort']})",
        "debt": f"={c}{R['rv_end']}+{c}{R['tl_end']}",
    }
    if variant == "growing":
        # The recurring link that goes missing: the PP&E roll adds capex but never subtracts D&A.
        f["ppe_end"] = f"=SUM({c}{R['ppe_beg']}:{c}{R['ppe_capex']})"
    return f


TOTALS = {"gp", "ebitda", "ebit", "ebt", "ni", "ppe_end", "ta", "tle", "cfo", "cfi", "cff",
          "netchg", "endcash", "rv_end", "tl_end", "debt"}
PCT_ROWS = {"growth", "gm", "sga", "da", "capex", "tax", "rdebt", "rcash", "payout"}

LABELS = {
    "toggle": "Circularity toggle (1 = on, 0 = off)",
    "growth": "Revenue growth", "gm": "Gross margin", "sga": "SG&A, % of revenue",
    "da": "D&A, % of revenue", "capex": "Capex, % of revenue",
    "dso": "DSO (days, on revenue)", "dio": "DIO (days, on COGS)", "dpo": "DPO (days, on COGS)",
    "tax": "Tax rate", "rdebt": "Interest rate on debt", "rcash": "Interest rate on cash",
    "amort": "Term loan amortization", "payout": "Dividend payout, % of net income",
    "mincash": "Minimum cash",
    "rev": "Revenue", "cogs": "Cost of goods sold", "gp": "Gross profit", "sgax": "SG&A",
    "ebitda": "EBITDA", "dax": "D&A", "ebit": "EBIT", "intexp": "Interest expense",
    "intinc": "Interest income", "ebt": "Pre-tax income", "taxes": "Taxes", "ni": "Net income",
    "div": "Dividends",
    "ppe_beg": "Beginning net PP&E", "ppe_capex": "Plus: capex", "ppe_da": "Less: D&A",
    "ppe_end": "Ending net PP&E",
    "wc_ar": "Accounts receivable", "wc_inv": "Inventory", "wc_ap": "Accounts payable",
    "d_ar": "(Increase) in receivables", "d_inv": "(Increase) in inventory",
    "d_ap": "Increase in payables",
    "cash": "Cash", "ar": "Accounts receivable", "inv": "Inventory", "ppe": "Net PP&E",
    "ta": "Total assets", "ap": "Accounts payable", "rev_bs": "Revolver", "tl": "Term loan",
    "cs": "Common stock", "re": "Retained earnings", "tle": "Total liabilities and equity",
    "check": "Balance check (assets minus liabilities and equity)",
    "cf_ni": "Net income", "cf_da": "Plus: D&A", "cf_ar": "Less: increase in receivables",
    "cf_inv": "Less: increase in inventory", "cf_ap": "Plus: increase in payables",
    "cfo": "Cash from operations", "cf_capex": "Capex", "cfi": "Cash from investing",
    "cf_amort": "Term loan amortization", "cf_draw": "Revolver draw",
    "cf_repay": "Revolver repayment", "cf_div": "Dividends", "cff": "Cash from financing",
    "netchg": "Net change in cash", "begcash": "Beginning cash", "endcash": "Ending cash",
    "cash_pre": "Memo: cash before revolver",
    "dr_min": "Minimum cash", "dr_pre": "Cash before revolver", "rv_beg": "Beginning revolver",
    "rv_draw": "Revolver draw", "rv_repay": "Revolver repayment", "rv_end": "Ending revolver",
    "tl_beg": "Beginning term loan", "tl_amort": "Less: amortization", "tl_end": "Ending term loan",
    "debt": "Total debt (revolver plus term loan)",
}
INDENTED = {"ppe_capex", "ppe_da", "tl_amort", "cf_da", "cf_ar", "cf_inv", "cf_ap", "cf_capex",
            "cf_amort", "cf_draw", "cf_repay", "cf_div", "rv_draw", "rv_repay"}

SECTIONS = [
    ("drv_hdr", "Drivers"), ("is_hdr", "Income statement"), ("ppe_hdr", "PP&E schedule"),
    ("wc_hdr", "Working capital schedule"), ("bs_hdr", "Balance sheet"),
    ("cf_hdr", "Cash flow statement"), ("dr_hdr", "Debt and revolver"),
]


def section(ws, key, text, formula_col):
    from openpyxl.styles import Alignment, Font
    row = R[key]
    last = 9 if formula_col else 7
    header_bar(ws, row, 2, last, text)
    for col, y in zip(COLS, YEARS):
        c = ws[f"{col}{row}"]
        c.value = y
        c.font = Font(name=XL_FONT, size=XL_SIZE, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="right")
    if formula_col:
        c = ws[f"{FCOL}{row}"]
        c.value = "Formula in the 2026E column"
        c.font = Font(name=XL_FONT, size=XL_SIZE, bold=True, color="FFFFFF")


def na(ws, ref):
    put(ws, ref, "n/a", kind="label", align="right", color="8A8984")


def label(ws, key, bold=False):
    put(ws, f"B{R[key]}", LABELS[key], bold=bold or key in TOTALS or key == "check",
        indent=1 if key in INDENTED else 0)


def write_model(ws, variant=None, formula_col=True):
    """Write the full model onto ws. variant: None, "typo" (constant gap) or "growing"."""
    from openpyxl.styles import Font
    set_widths(ws, {"A": 2, "B": 46, "C": 2, "D": 11, "E": 11, "F": 11, "G": 11, "H": 2, "I": 38})
    ws.sheet_view.showGridLines = True

    put(ws, f"B{R['title']}", "Three statement model, modeling test build", bold=True, size=12)
    put(ws, f"B{R['units']}", "$ millions, fiscal years ending December 31", italic=True, color="52514E")
    year_header(ws, R["units"], 4, YEARS)

    for key, text in SECTIONS:
        section(ws, key, text, formula_col)

    # ---- drivers
    label(ws, "toggle")
    put(ws, f"D{R['toggle']}", TOGGLE_DEFAULT, kind="input", fmt=FMT_TOGGLE)
    if formula_col:
        put(ws, f"{FCOL}{R['toggle']}", "Input. 1 = average debt, 0 = beginning debt",
            color="52514E")
    for key in ["growth", "gm", "sga", "da", "capex", "dso", "dio", "dpo", "tax", "rdebt", "rcash",
                "amort", "payout", "mincash"]:
        label(ws, key)
    # historical implied ratios (formulas off the 2025A statements)
    hist_ratio = {
        "gm": f"=D{R['gp']}/D{R['rev']}", "sga": f"=-D{R['sgax']}/D{R['rev']}",
        "da": f"=-D{R['dax']}/D{R['rev']}", "capex": f"=D{R['ppe_capex']}/D{R['rev']}",
        "dso": f"=D{R['wc_ar']}/D{R['rev']}*365", "dio": f"=-D{R['wc_inv']}/D{R['cogs']}*365",
        "dpo": f"=-D{R['wc_ap']}/D{R['cogs']}*365", "tax": f"=-D{R['taxes']}/D{R['ebt']}",
        "rdebt": f"=-D{R['intexp']}/D{R['debt']}", "payout": f"=-D{R['div']}/D{R['ni']}",
    }
    for key in ["growth", "rcash", "amort", "mincash"]:
        na(ws, f"D{R[key]}")
    for key, fml in hist_ratio.items():
        fmt = FMT_PCT if key in PCT_ROWS else FMT_DAYS_HIST
        put(ws, f"D{R[key]}", fml, kind="formula", fmt=fmt)
    for key, vals in DRV.items():
        for col, v in zip(COLS[1:], vals):
            if key in PCT_ROWS:
                fmt = FMT_PCT
            elif key in ("dso", "dio", "dpo"):
                fmt = FMT_DAYS
            else:
                fmt = FMT_NUM
            put(ws, f"{col}{R[key]}", v, kind="input", fmt=fmt)

    # ---- income statement, schedules, balance sheet, cash flow, debt: labels
    for key in ["rev", "cogs", "gp", "sgax", "ebitda", "dax", "ebit", "intexp", "intinc", "ebt",
                "taxes", "ni", "div", "ppe_beg", "ppe_capex", "ppe_da", "ppe_end", "wc_ar", "wc_inv",
                "wc_ap", "d_ar", "d_inv", "d_ap", "cash", "ar", "inv", "ppe", "ta", "ap", "rev_bs",
                "tl", "cs", "re", "tle", "check", "cf_ni", "cf_da", "cf_ar", "cf_inv", "cf_ap", "cfo",
                "cf_capex", "cfi", "cf_amort", "cf_draw", "cf_repay", "cf_div", "cff", "netchg",
                "begcash", "endcash", "cash_pre", "dr_min", "dr_pre", "rv_beg", "rv_draw", "rv_repay",
                "rv_end", "tl_beg", "tl_amort", "tl_end", "debt"]:
        label(ws, key)

    # ---- 2025A column: typed historicals in blue, subtotals as formulas
    hist_bs = dict(HIST_BS)
    if variant == "typo":
        hist_bs["ppe"] = 412.0   # the mistyped historical: filing says 400.0
    for key, v in HIST_IS.items():
        put(ws, f"D{R[key]}", v, kind="input", fmt=FMT_NUM)
    for key, v in hist_bs.items():
        put(ws, f"D{R[key]}", v, kind="input", fmt=FMT_NUM)
    put(ws, f"D{R['ppe_capex']}", HIST_CAPEX, kind="input", fmt=FMT_NUM)
    d_formulas = {
        "gp": f"=SUM(D{R['rev']}:D{R['cogs']})", "ebitda": f"=SUM(D{R['gp']}:D{R['sgax']})",
        "ebit": f"=SUM(D{R['ebitda']}:D{R['dax']})", "ebt": f"=SUM(D{R['ebit']}:D{R['intinc']})",
        "ni": f"=SUM(D{R['ebt']}:D{R['taxes']})",
        "ppe_da": f"=D{R['dax']}", "ppe_end": f"=D{R['ppe']}",
        "wc_ar": f"=D{R['ar']}", "wc_inv": f"=D{R['inv']}", "wc_ap": f"=D{R['ap']}",
        "ta": f"=SUM(D{R['cash']}:D{R['ppe']})", "tle": f"=SUM(D{R['ap']}:D{R['re']})",
        "check": f"=D{R['ta']}-D{R['tle']}",
        "rv_end": f"=D{R['rev_bs']}", "tl_end": f"=D{R['tl']}", "debt": f"=D{R['rv_end']}+D{R['tl_end']}",
    }
    for key, fml in d_formulas.items():
        kind = "total" if key in TOTALS else "formula"
        put(ws, f"D{R[key]}", fml, kind=kind, fmt=FMT_CHECK if key == "check" else FMT_NUM)

    # ---- projected columns
    for i, c in enumerate(COLS[1:]):
        p = COLS[i]
        for key, fml in projected_formulas(c, p, variant).items():
            kind = "total" if key in TOTALS else "formula"
            put(ws, f"{c}{R[key]}", fml, kind=kind, fmt=FMT_CHECK if key == "check" else FMT_NUM)
            if formula_col and c == "E":
                cell = ws[f"{FCOL}{R[key]}"]
                cell.value = fml
                cell.data_type = "s"
                cell.font = Font(name=XL_FONT, size=XL_SIZE, color="52514E")

    # ---- check row styling: subtle fill, red in the broken variants
    for col in ["B"] + COLS:
        c = ws[f"{col}{R['check']}"]
        from openpyxl.styles import PatternFill
        c.fill = PatternFill("solid", fgColor=FILL_SUBTLE)
        if variant:
            c.font = Font(name=XL_FONT, size=XL_SIZE, bold=True, color=RED)
    if variant:
        ws.column_dimensions[FCOL].width = 40
        # extend the balance sheet bar over the note column so the callouts have a header
        header_bar(ws, R["bs_hdr"], 8, 9, "")
        ws[f"B{R['bs_hdr']}"].value = "Balance sheet"
        put(ws, f"{FCOL}{R['bs_hdr']}", "What went wrong", kind="header")
    if variant == "typo":
        put(ws, f"{FCOL}{R['ppe']}", "Typed 412.0, the filing says 400.0", color=RED, bold=True)
        put(ws, f"{FCOL}{R['check']}", "Same 12.0 every year: a one time break", color=RED)
    elif variant == "growing":
        put(ws, f"{FCOL}{R['ppe']}", "PP&E roll never subtracts D&A", color=RED, bold=True)
        put(ws, f"{FCOL}{R['check']}", "Grows by each year's D&A: a missing link", color=RED)
    ws.freeze_panes = None


def build(out=OUT):
    wb = new_workbook()
    wb.properties.creator = "Wall Street Guide"
    wb.properties.lastModifiedBy = "Wall Street Guide"
    ws = wb.active
    ws.title = "Model"
    write_model(ws, variant=None, formula_col=True)
    ws2 = wb.create_sheet("Error constant")
    write_model(ws2, variant="typo", formula_col=False)
    ws3 = wb.create_sheet("Error growing")
    write_model(ws3, variant="growing", formula_col=False)
    # Let Excel iterate if a reader flips the toggle to 1 (the circular, average balance case).
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(iterate=True, iterateCount=100, iterateDelta=0.001, fullCalcOnLoad=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ---------------------------------------------------------------- article tables (section 9) for the check
ARTICLE = {
    "Model": {
        # income statement
        "rev": [1000.0, 1100.0, 1188.0, 1259.3], "cogs": [-600.0, -660.0, -712.8, -755.6],
        "gp": [400.0, 440.0, 475.2, 503.7], "sgax": [-250.0, -275.0, -297.0, -314.8],
        "ebitda": [150.0, 165.0, 178.2, 188.9], "dax": [-40.0, -44.0, -47.5, -50.4],
        "ebit": [110.0, 121.0, 130.7, 138.5], "intexp": [-18.0, -18.0, -16.2, -16.4],
        "intinc": [2.0, 1.6, 1.7, 1.0], "ebt": [94.0, 104.6, 116.2, 123.1],
        "taxes": [-18.8, -20.9, -23.2, -24.6], "ni": [75.2, 83.7, 93.0, 98.5],
        "div": [-18.8, -20.9, -23.2, -24.6],
        # balance sheet
        "cash": [80.0, 86.9, 50.0, 50.0], "ar": [125.0, 135.6, 146.5, 155.3],
        "inv": [100.0, 108.5, 117.2, 124.2], "ppe": [400.0, 411.0, 506.0, 518.6],
        "ta": [705.0, 742.0, 819.7, 848.1], "ap": [50.0, 54.2, 58.6, 62.1],
        "rev_bs": [0.0, 0.0, 33.6, 14.6], "tl": [300.0, 270.0, 240.0, 210.0],
        "cs": [100.0, 100.0, 100.0, 100.0], "re": [255.0, 317.8, 387.5, 461.4],
        "tle": [705.0, 742.0, 819.7, 848.1], "check": [0.0, 0.0, 0.0, 0.0],
        # cash flow statement (2026E to 2028E)
        "cf_ni": [None, 83.7, 93.0, 98.5], "cf_da": [None, 44.0, 47.5, 50.4],
        "cf_ar": [None, -10.6, -10.8, -8.8], "cf_inv": [None, -8.5, -8.7, -7.0],
        "cf_ap": [None, 4.2, 4.3, 3.5], "cfo": [None, 112.8, 125.3, 136.6],
        "cf_capex": [None, -55.0, -142.6, -63.0], "cfi": [None, -55.0, -142.6, -63.0],
        "cf_amort": [None, -30.0, -30.0, -30.0], "cf_draw": [None, 0.0, 33.6, 0.0],
        "cf_repay": [None, 0.0, 0.0, -19.0], "cf_div": [None, -20.9, -23.2, -24.6],
        "cff": [None, -50.9, -19.6, -73.6], "netchg": [None, 6.9, -36.9, 0.0],
        "cash_pre": [None, 86.9, 16.4, 69.0],
        # drivers, historical implied
        "gm": [0.40, None, None, None], "sga": [0.25, None, None, None], "da": [0.04, None, None, None],
        "capex": [0.05, None, None, None], "dso": [45.6, None, None, None], "dio": [60.8, None, None, None],
        "dpo": [30.4, None, None, None], "tax": [0.20, None, None, None], "rdebt": [0.06, None, None, None],
        "payout": [0.25, None, None, None],
    },
    "Error constant": {"check": [12.0, 12.0, 12.0, 12.0]},
    "Error growing": {"check": [0.0, 44.0, 91.5, 141.9]},
}


def check_against_article(xlsx=OUT, verbose=True):
    """Recalculate with LibreOffice and compare every article table cell (one decimal)."""
    vals = recalc_values(xlsx)
    mismatches = []
    n = 0
    for sheet, rows in ARTICLE.items():
        for key, expected in rows.items():
            for col, exp in zip(COLS, expected):
                if exp is None:
                    continue
                got = vals[sheet].get(f"{col}{R[key]}")
                n += 1
                if got is None or not isinstance(got, (int, float)):
                    mismatches.append((sheet, key, col, exp, got))
                    continue
                places = 3 if key in PCT_ROWS else 1
                if round(got + 0.0, places) != round(exp, places):
                    mismatches.append((sheet, key, col, exp, got))
    checks = {s: [vals[s].get(f"{c}{R['check']}") for c in COLS] for s in ARTICLE}
    if verbose:
        print(f"compared {n} cells against the article tables; mismatches: {len(mismatches)}")
        for m in mismatches:
            print("  MISMATCH", m)
        for s, v in checks.items():
            print(f"  check row {s}: {v}")
    return mismatches, checks, vals


if __name__ == "__main__":
    path = build()
    print("built", path)
    mm, checks, _ = check_against_article(path)
    assert not mm, "workbook disagrees with the article"
    assert all(abs(v) < 1e-9 for v in checks["Model"]), "balance check is not zero"
