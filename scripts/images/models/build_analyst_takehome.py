#!/usr/bin/env python3
"""Build scripts/images/models/analyst-takehome.xlsx: the gross-to-net tax computation behind
"Investment Banking Analyst Salary 2027: What You Actually Take Home" (October 5, 2026).

One first-year analyst in New York City, single filer, $120,000 base plus an $80,000 bonus, no
pre-tax deductions, on the 2026 rules the article uses. Every parameter is a blue input and every
tax line is a live formula, so the screenshots show the bracket math working:

  Sheet "Take-home"
    pay and federal taxable income (gross less the $16,100 standard deduction)
    federal income tax by bracket, all seven 2026 single-filer slices, MIN / MAX slice formulas
    Social Security (capped at the $184,500 wage base), Medicare and the 0.9 percent additional tax
    New York State: the $8,000 deduction, the published schedule for reference, and the flat
      5.9 percent that worksheet 7 applies once New York adjusted gross income reaches $157,650
    New York City resident tax by bracket, four slices
    the take-home block: total taxes, take-home, share of gross, effective rate, per month
  Sheet "NY schedule"
    the nine-step New York State rate schedule applied slice by slice, for the reference line

The inputs are read from scripts/images/models/numbers_investment-banking-analyst-salary-2027.py,
the plain-Python source of every figure in the article, so the workbook can never drift from it.

Run from the repo root:  python3 scripts/images/models/build_analyst_takehome.py
Idempotent: rewrites the workbook every time. Workbook author: Wall Street Guide.
"""
import importlib.util
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts" / "images"))
from wsg_viz import XL_FONT, XL_SIZE, header_bar, new_workbook, put, set_widths  # noqa: E402

OUT_XLSX = ROOT / "scripts" / "images" / "models" / "analyst-takehome.xlsx"
NUMBERS = ROOT / "scripts" / "images" / "models" / "numbers_investment-banking-analyst-salary-2027.py"

FMT_INT = '#,##0;(#,##0);"-"'        # dash for an untouched bracket
FMT_INT_Z = '#,##0;(#,##0);0'        # zero shown as 0 (the article prints $0)
FMT_PCT1 = '0.0%'
FMT_PCT2 = '0.00%'
FMT_PCT3 = '0.000%'
FMT_DOLLAR0 = '"$"#,##0;("$"#,##0);"$"0'


def load_numbers():
    """Import the article's numbers script (its file name has hyphens, so go through importlib)."""
    spec = importlib.util.spec_from_file_location("salary_numbers", NUMBERS)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def col_headers(ws, row, items):
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin", color="000000")
    for col, text in items.items():
        c = ws[f"{col}{row}"]
        c.value = text
        c.font = Font(name=XL_FONT, size=XL_SIZE, bold=True)
        c.alignment = Alignment(horizontal="right")
        c.border = Border(bottom=thin)


def right(ws, ref):
    from openpyxl.styles import Alignment
    ws[ref].alignment = Alignment(horizontal="right")


def build(out_path=OUT_XLSX):
    N = load_numbers()
    base, bonus = 120_000, 80_000

    wb = new_workbook()
    wb.properties.creator = "Wall Street Guide"
    wb.properties.lastModifiedBy = "Wall Street Guide"
    wb.properties.title = "First-year analyst take-home, New York City"
    ws = wb.active
    ws.title = "Take-home"
    set_widths(ws, {"A": 2, "B": 52, "C": 2, "D": 12, "E": 12, "F": 9, "G": 15, "H": 13})
    ws.sheet_view.showGridLines = False
    R = {}

    put(ws, "B2", "First-year analyst in New York City: gross pay to take-home", bold=True, size=12)
    put(ws, "B3", "Single filer, 2026 rules, no pre-tax deductions. Blue on yellow = typed input, black = formula, green = pulled from another sheet.", italic=True)

    # ---- pay and federal taxable income
    r = 5
    R["pay_header"] = r
    header_bar(ws, r, 2, 8, "PAY AND FEDERAL TAXABLE INCOME"); r += 1
    R["base"] = r; put(ws, f"B{r}", "Base salary"); put(ws, f"H{r}", base, kind="input", fmt=FMT_INT_Z); r += 1
    R["bonus"] = r; put(ws, f"B{r}", "Year-end bonus"); put(ws, f"H{r}", bonus, kind="input", fmt=FMT_INT_Z); r += 1
    R["gross"] = r; put(ws, f"B{r}", "Gross pay", kind="total"); put(ws, f"H{r}", f"=H{R['base']}+H{R['bonus']}", kind="total", fmt=FMT_INT_Z); r += 1
    R["std"] = r; put(ws, f"B{r}", "Less: federal standard deduction (single, 2026)"); put(ws, f"H{r}", N.FED_STD_DED, kind="input", fmt=FMT_INT_Z); r += 1
    R["fed_taxable"] = r; put(ws, f"B{r}", "Federal taxable income", kind="total"); put(ws, f"H{r}", f"=H{R['gross']}-H{R['std']}", kind="total", fmt=FMT_INT_Z); r += 2

    # ---- federal brackets
    R["fed_header"] = r
    header_bar(ws, r, 2, 8, "FEDERAL INCOME TAX BY BRACKET (2026, single filer)"); r += 1
    col_headers(ws, r, {"D": "Slice from", "E": "Slice to", "F": "Rate", "G": "Dollars in slice", "H": "Tax"}); r += 1
    R["fed_first"] = r
    prev_top = None
    for i, (top, rate) in enumerate(N.FED_BRACKETS):
        put(ws, f"B{r}", f"Bracket {i + 1}")
        if i == 0:
            put(ws, f"D{r}", 0, kind="formula", fmt=FMT_INT_Z)
        else:
            put(ws, f"D{r}", f"=E{r - 1}", kind="formula", fmt=FMT_INT_Z)
        if top == float("inf"):
            put(ws, f"E{r}", "and up", align="right")
            put(ws, f"G{r}", f"=MAX(0,$H${R['fed_taxable']}-D{r})", kind="formula", fmt=FMT_INT)
        else:
            put(ws, f"E{r}", top, kind="input", fmt=FMT_INT_Z)
            put(ws, f"G{r}", f"=MAX(0,MIN($H${R['fed_taxable']},E{r})-D{r})", kind="formula", fmt=FMT_INT)
        put(ws, f"F{r}", rate, kind="input", fmt=FMT_PCT1)
        put(ws, f"H{r}", f"=G{r}*F{r}", kind="formula", fmt=FMT_INT)
        r += 1
    R["fed_last"] = r - 1
    R["fed_total"] = r
    put(ws, f"B{r}", "Total federal income tax", kind="total")
    put(ws, f"G{r}", f"=SUM(G{R['fed_first']}:G{R['fed_last']})", kind="total", fmt=FMT_INT_Z)
    put(ws, f"H{r}", f"=SUM(H{R['fed_first']}:H{R['fed_last']})", kind="total", fmt=FMT_INT_Z); r += 1
    R["fed_eff"] = r
    put(ws, f"B{r}", "Effective federal rate on gross pay (the top rate touched is 24%)", italic=True)
    put(ws, f"H{r}", f"=H{R['fed_total']}/H{R['gross']}", kind="formula", fmt=FMT_PCT1, italic=True); r += 2

    # ---- FICA
    R["fica_header"] = r
    header_bar(ws, r, 2, 8, "SOCIAL SECURITY AND MEDICARE"); r += 1
    col_headers(ws, r, {"F": "Rate", "G": "Wages taxed", "H": "Tax"}); r += 1
    R["wage_base"] = r; put(ws, f"B{r}", "Social Security wage base (2026)"); put(ws, f"G{r}", N.SS_WAGE_BASE, kind="input", fmt=FMT_INT_Z); r += 1
    R["ss"] = r
    put(ws, f"B{r}", "Social Security, 6.2% of wages up to the wage base")
    put(ws, f"F{r}", N.SS_RATE, kind="input", fmt=FMT_PCT1)
    put(ws, f"G{r}", f"=MIN(H{R['gross']},G{R['wage_base']})", kind="formula", fmt=FMT_INT_Z)
    put(ws, f"H{r}", f"=F{r}*G{r}", kind="formula", fmt=FMT_INT_Z); r += 1
    R["med"] = r
    put(ws, f"B{r}", "Medicare, 1.45% of all wages, no cap")
    put(ws, f"F{r}", N.MEDICARE_RATE, kind="input", fmt=FMT_PCT2)
    put(ws, f"G{r}", f"=H{R['gross']}", kind="formula", fmt=FMT_INT_Z)
    put(ws, f"H{r}", f"=F{r}*G{r}", kind="formula", fmt=FMT_INT_Z); r += 1
    R["addl"] = r
    put(ws, f"B{r}", "Additional Medicare, 0.9% above the threshold in column D")
    put(ws, f"D{r}", N.ADDL_MEDICARE_THRESHOLD, kind="input", fmt=FMT_INT_Z)
    put(ws, f"F{r}", N.ADDL_MEDICARE_RATE, kind="input", fmt=FMT_PCT1)
    put(ws, f"G{r}", f"=MAX(0,H{R['gross']}-D{r})", kind="formula", fmt=FMT_INT_Z)
    put(ws, f"H{r}", f"=F{r}*G{r}", kind="formula", fmt=FMT_INT_Z); r += 1
    R["fica_total"] = r
    put(ws, f"B{r}", "Total Social Security and Medicare", kind="total")
    put(ws, f"H{r}", f"=SUM(H{R['ss']}:H{R['addl']})", kind="total", fmt=FMT_INT_Z); r += 2
    R["fica_end"] = r - 2

    # ---- New York State
    R["nys_header"] = r
    header_bar(ws, r, 2, 8, "NEW YORK STATE (single filer, 2026)"); r += 1
    col_headers(ws, r, {"F": "Rate", "G": "Taxed on", "H": "Tax"}); r += 1
    R["nyagi"] = r; put(ws, f"B{r}", "New York adjusted gross income (wages, no deferrals)"); put(ws, f"H{r}", f"=H{R['gross']}", kind="formula", fmt=FMT_INT_Z); r += 1
    R["ny_std"] = r; put(ws, f"B{r}", "Less: New York standard deduction (single)"); put(ws, f"H{r}", N.NYS_STD_DED, kind="input", fmt=FMT_INT_Z); r += 1
    R["ny_taxable"] = r; put(ws, f"B{r}", "New York taxable income", kind="total"); put(ws, f"H{r}", f"=H{R['nyagi']}-H{R['ny_std']}", kind="total", fmt=FMT_INT_Z); r += 1
    R["ws7_threshold"] = r
    put(ws, f"B{r}", "Worksheet 7 applies once New York AGI reaches (phase-out complete)")
    put(ws, f"H{r}", N.NYS_RECAPTURE_START + N.NYS_RECAPTURE_WIDTH, kind="input", fmt=FMT_INT_Z); r += 1
    R["ny_sched"] = r
    put(ws, f"B{r}", "Tax from the published rate schedule, for reference (NY schedule sheet)")
    put(ws, f"H{r}", "='NY schedule'!H15", kind="link", fmt=FMT_INT_Z); r += 1
    R["ny_flat"] = r
    put(ws, f"B{r}", "Tax at the flat rate on the whole taxable income")
    put(ws, f"F{r}", N.NYS_FLAT_RATE_1, kind="input", fmt=FMT_PCT1)
    put(ws, f"G{r}", f"=H{R['ny_taxable']}", kind="formula", fmt=FMT_INT_Z)
    put(ws, f"H{r}", f"=F{r}*G{r}", kind="formula", fmt=FMT_INT_Z); r += 1
    R["nys_tax"] = r
    put(ws, f"B{r}", "New York State tax (flat rate if AGI is at or above the threshold)", kind="total")
    put(ws, f"H{r}", f"=IF(H{R['nyagi']}>=H{R['ws7_threshold']},H{R['ny_flat']},H{R['ny_sched']})", kind="total", fmt=FMT_INT_Z); r += 1
    R["ny_benefit"] = r
    put(ws, f"B{r}", "Lower-bracket benefit the worksheet takes away (flat minus schedule)", italic=True)
    put(ws, f"H{r}", f"=H{R['ny_flat']}-H{R['ny_sched']}", kind="formula", fmt=FMT_INT_Z, italic=True); r += 2

    # ---- New York City
    R["nyc_header"] = r
    header_bar(ws, r, 2, 8, "NEW YORK CITY RESIDENT TAX BY BRACKET (2026, single filer)"); r += 1
    col_headers(ws, r, {"D": "Slice from", "E": "Slice to", "F": "Rate", "G": "Dollars in slice", "H": "Tax"}); r += 1
    R["nyc_first"] = r
    for i, (top, rate) in enumerate(N.NYC_BRACKETS):
        put(ws, f"B{r}", f"Bracket {i + 1}")
        if i == 0:
            put(ws, f"D{r}", 0, kind="formula", fmt=FMT_INT_Z)
        else:
            put(ws, f"D{r}", f"=E{r - 1}", kind="formula", fmt=FMT_INT_Z)
        if top == float("inf"):
            put(ws, f"E{r}", "and up", align="right")
            put(ws, f"G{r}", f"=MAX(0,$H${R['ny_taxable']}-D{r})", kind="formula", fmt=FMT_INT)
        else:
            put(ws, f"E{r}", top, kind="input", fmt=FMT_INT_Z)
            put(ws, f"G{r}", f"=MAX(0,MIN($H${R['ny_taxable']},E{r})-D{r})", kind="formula", fmt=FMT_INT)
        put(ws, f"F{r}", rate, kind="input", fmt=FMT_PCT3)
        put(ws, f"H{r}", f"=G{r}*F{r}", kind="formula", fmt=FMT_INT)
        r += 1
    R["nyc_last"] = r - 1
    R["nyc_total"] = r
    put(ws, f"B{r}", "Total New York City tax", kind="total")
    put(ws, f"G{r}", f"=SUM(G{R['nyc_first']}:G{R['nyc_last']})", kind="total", fmt=FMT_INT_Z)
    put(ws, f"H{r}", f"=SUM(H{R['nyc_first']}:H{R['nyc_last']})", kind="total", fmt=FMT_INT_Z); r += 2

    # ---- take-home
    R["th_header"] = r
    header_bar(ws, r, 2, 8, "TAKE-HOME"); r += 1
    R["total_tax"] = r
    put(ws, f"B{r}", "Total taxes (federal, Social Security and Medicare, state, city)", kind="total")
    put(ws, f"H{r}", f"=H{R['fed_total']}+H{R['fica_total']}+H{R['nys_tax']}+H{R['nyc_total']}", kind="total", fmt=FMT_INT_Z); r += 1
    R["net"] = r
    put(ws, f"B{r}", "Take-home pay", bold=True)
    put(ws, f"H{r}", f"=H{R['gross']}-H{R['total_tax']}", kind="formula", fmt=FMT_INT_Z, bold=True); r += 1
    R["net_share"] = r
    put(ws, f"B{r}", "Take-home as a share of gross")
    put(ws, f"H{r}", f"=H{R['net']}/H{R['gross']}", kind="formula", fmt=FMT_PCT1); r += 1
    R["eff"] = r
    put(ws, f"B{r}", "Effective tax rate")
    put(ws, f"H{r}", f"=H{R['total_tax']}/H{R['gross']}", kind="formula", fmt=FMT_PCT1); r += 1
    R["monthly"] = r
    put(ws, f"B{r}", "Take-home per month")
    put(ws, f"H{r}", f"=H{R['net']}/12", kind="formula", fmt=FMT_INT_Z); r += 1
    R["th_end"] = r - 1

    # ================================================================ NY schedule sheet
    ns = wb.create_sheet("NY schedule")
    set_widths(ns, {"A": 2, "B": 44, "C": 2, "D": 13, "E": 13, "F": 9, "G": 15, "H": 13})
    ns.sheet_view.showGridLines = False
    put(ns, "B2", "New York State published rate schedule, single filer, 2026 (for reference)", bold=True, size=12)
    put(ns, "B3", "What the schedule would charge on the analyst's taxable income before worksheet 7 phases the lower brackets out.", italic=True)
    put(ns, "B4", "New York taxable income (from the Take-home sheet)")
    put(ns, "H4", f"='Take-home'!H{R['ny_taxable']}", kind="link", fmt=FMT_INT_Z)
    col_headers(ns, 5, {"D": "Slice from", "E": "Slice to", "F": "Rate", "G": "Dollars in slice", "H": "Tax"})
    rr = 6
    for i, (top, rate) in enumerate(N.NYS_BRACKETS):
        put(ns, f"B{rr}", f"Bracket {i + 1}")
        put(ns, f"D{rr}", 0 if i == 0 else f"=E{rr - 1}", kind="formula", fmt=FMT_INT_Z)
        if top == float("inf"):
            put(ns, f"E{rr}", "and up", align="right")
            put(ns, f"G{rr}", f"=MAX(0,$H$4-D{rr})", kind="formula", fmt=FMT_INT)
        else:
            put(ns, f"E{rr}", top, kind="input", fmt=FMT_INT_Z)
            put(ns, f"G{rr}", f"=MAX(0,MIN($H$4,E{rr})-D{rr})", kind="formula", fmt=FMT_INT)
        put(ns, f"F{rr}", rate, kind="input", fmt=FMT_PCT2)
        put(ns, f"H{rr}", f"=G{rr}*F{rr}", kind="formula", fmt=FMT_INT)
        rr += 1
    assert rr == 15, rr
    put(ns, "B15", "Tax from the rate schedule", kind="total")
    put(ns, "G15", "=SUM(G6:G14)", kind="total", fmt=FMT_INT_Z)
    put(ns, "H15", "=SUM(H6:H14)", kind="total", fmt=FMT_INT_Z)

    wb.calculation.fullCalcOnLoad = True
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, R


if __name__ == "__main__":
    path, R = build()
    print(f"wrote {path}")
    print("rows:", R)
